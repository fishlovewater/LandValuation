"""User-facing XLSX export for a finalized valuation report package.

This workbook is a structured data export of the finalized report package.  It
is deliberately not presented as an original government Excel template and it
never exposes internal UUIDs, MinIO paths, engineering factor codes, or storage
metadata.
"""

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.valuation.report_packages.factor_catalog import (
    INDIVIDUAL_FACTOR_BY_CODE,
    TEMPLATE_FACTOR_BY_CODE,
)
from app.valuation.rule_packs.coverage import NEW_TAIPEI_DISTRICTS


XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_SUMMARY = "案件摘要"
SHEET_S01 = "地價區段勘查表"
SHEET_F02_RF = "區域因素分析"
SHEET_F02 = "比較法調查估價表"

_LAND_USE_LABELS = {
    "RESIDENTIAL": "住宅用地",
    "COMMERCIAL": "商業用地",
    "INDUSTRIAL": "工業用地",
    "AGRICULTURAL": "農業用地",
    "OTHER": "其他用途",
}

_THIN = Side(style="thin", color="D9E1E8")
_HEADER_FILL = PatternFill("solid", fgColor="EAF1F8")
_SECTION_FILL = PatternFill("solid", fgColor="F5F7FA")


def _display(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    return value


def _factor_label(code: str | None, *, individual: bool = False) -> str:
    catalog = INDIVIDUAL_FACTOR_BY_CODE if individual else TEMPLATE_FACTOR_BY_CODE
    item = catalog.get(code or "")
    return item.label if item is not None else "其他影響因素"


def _district_label(code: str | None) -> str:
    if not code:
        return ""
    return NEW_TAIPEI_DISTRICTS.get(code, code)


def _land_use_label(value: str | None) -> str:
    if not value:
        return ""
    return _LAND_USE_LABELS.get(value.upper(), value)


def _style_table(sheet, header_row: int, *, widths: list[int] | None = None) -> None:
    for cell in sheet[header_row]:
        if cell.value is None:
            continue
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=_THIN)
    sheet.freeze_panes = f"A{header_row + 1}"
    if widths:
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _append_section(sheet, title: str) -> None:
    row = sheet.max_row + 1
    sheet.append([title])
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    cell = sheet.cell(row=row, column=1)
    cell.font = Font(bold=True)
    cell.fill = _SECTION_FILL
    cell.alignment = Alignment(vertical="center")


def build_formal_report_xlsx(data: dict[str, Any]) -> bytes:
    """Render a finalized report package into a readable multi-sheet workbook."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    _build_summary(workbook.create_sheet(SHEET_SUMMARY), data)
    _build_s01(workbook.create_sheet(SHEET_S01), data)
    _build_f02_rf(workbook.create_sheet(SHEET_F02_RF), data)
    _build_f02(workbook.create_sheet(SHEET_F02), data)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _build_summary(sheet, data: dict[str, Any]) -> None:
    context = data.get("context") or {}
    sheet.append(["正式查估資料 Excel", ""])
    sheet.merge_cells("A1:B1")
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.append([
        "說明",
        "本檔為系統依已完成正式檢核之查估資料產生的結構化 Excel；正式送審主要文件仍為完整送審 PDF。",
    ])
    sheet.append([])
    sheet.append(["項目", "內容"])
    rows = [
        ("案件編號", context.get("case_no")),
        ("案件名稱", context.get("case_title")),
        ("行政區", _district_label(context.get("district_code"))),
        ("估價基準日", context.get("valuation_base_date")),
        ("土地用途", _land_use_label(context.get("land_use_type"))),
        ("查估資料版本", f"第 {data.get('version_no')} 版" if data.get("version_no") else ""),
    ]
    for label, value in rows:
        sheet.append([label, _display(value)])

    _append_section(sheet, "宗地資料")
    parcel_header = sheet.max_row + 1
    sheet.append(["行政區", "地段", "小段", "地號", "面積（平方公尺）", "使用分區", "編定用途"])
    for parcel in context.get("parcels") or []:
        sheet.append([
            _district_label(parcel.get("district_code")),
            _display(parcel.get("section_name")),
            _display(parcel.get("subsection_name")),
            _display(parcel.get("land_no")),
            _display(parcel.get("area_sqm")),
            _display(parcel.get("land_use_zone")),
            _display(parcel.get("designated_use")),
        ])

    benchmark_header = sheet.max_row + 2
    sheet.append([])
    sheet.append(["比準地", "地價區段"])
    for benchmark in context.get("benchmark_lands") or []:
        sheet.append([
            _display(benchmark.get("benchmark_land_no")),
            _display(benchmark.get("price_zone_no")),
        ])

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[4]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for cell in sheet[parcel_header]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for cell in sheet[benchmark_header]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    widths = [20, 46, 18, 18, 22, 20, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A5"


def _build_s01(sheet, data: dict[str, Any]) -> None:
    s01 = data.get("s01") or {}
    sheet.append(["地價區段勘查資料", ""])
    sheet.merge_cells("A1:B1")
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append(["項目", "內容"])
    base_rows = [
        ("地價區段", s01.get("district_name")),
        ("區段範圍", s01.get("district_boundary")),
        ("勘查日期", s01.get("survey_date")),
        ("都市計畫", s01.get("urban_plan_status")),
        ("使用分區", s01.get("land_use_zone")),
        ("建蔽率", s01.get("building_coverage_rate")),
        ("容積率", s01.get("floor_area_ratio")),
        ("禁止建築", s01.get("prohibited_building")),
        ("限制建築", s01.get("restricted_building")),
        ("主要道路", s01.get("main_road_name")),
        ("主要道路寬度（公尺）", s01.get("main_road_width_m")),
        ("區段內道路平均寬度（公尺）", s01.get("average_road_width_m")),
        ("現勘意見", s01.get("site_opinion")),
        ("備註", s01.get("notes")),
        ("承辦人", s01.get("handler_name")),
        ("課長", s01.get("section_head_name")),
        ("主管", s01.get("director_name")),
        ("估價人員", s01.get("appraiser_name")),
    ]
    for label, value in base_rows:
        sheet.append([label, _display(value)])

    sheet.append([])
    observation_header = sheet.max_row + 1
    sheet.append(["勘查項目", "填報內容", "設施名稱", "步行距離（公尺）", "來源說明", "人工確認"])
    for observation in s01.get("observations") or []:
        sheet.append([
            _factor_label(observation.get("item_code")),
            _display(observation.get("raw_value")),
            _display(observation.get("facility_name")),
            _display(observation.get("walking_distance_m")),
            _display(observation.get("source_notes")),
            _display(observation.get("confirmed_by_user")),
        ])
    _style_table(sheet, observation_header, widths=[30, 24, 24, 18, 44, 12])
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL


def _build_f02_rf(sheet, data: dict[str, Any]) -> None:
    regional = data.get("f02_rf") or {}
    sheet.append(["影響地價區域因素分析", ""])
    sheet.merge_cells("A1:B1")
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append(["項目", "內容"])
    sheet.append(["其他影響因素", _display(regional.get("other_influences"))])
    sheet.append(["備註", _display(regional.get("notes"))])
    sheet.append(["估價人員", _display(regional.get("appraiser_name"))])
    sheet.append(["計算狀態", "已完成" if regional.get("calculation_status") == "CALCULATED" else "尚未完成"])
    sheet.append(["計算時間", _display(regional.get("calculated_at"))])
    sheet.append([])

    header_row = sheet.max_row + 1
    sheet.append([
        "區域因素",
        "比準地填報等級",
        "比準地確認等級",
        "比較標的順序",
        "比較標的填報等級",
        "比較標的確認等級",
        "調整率",
        "來源說明",
        "人工確認",
    ])
    for factor in regional.get("factor_rows") or []:
        targets = factor.get("targets") or []
        if not targets:
            sheet.append([
                _factor_label(factor.get("factor_code")),
                _display(factor.get("benchmark_reported_level")),
                _display(factor.get("benchmark_confirmed_level")),
                "",
                "",
                "",
                "",
                _display(factor.get("source_notes")),
                _display(factor.get("confirmed_by_user")),
            ])
            continue
        for target in targets:
            sheet.append([
                _factor_label(factor.get("factor_code")),
                _display(factor.get("benchmark_reported_level")),
                _display(factor.get("benchmark_confirmed_level")),
                _display(target.get("display_order")),
                _display(target.get("reported_level")),
                _display(target.get("confirmed_level")),
                _display(target.get("calculated_adjustment_rate")),
                _display(target.get("source_notes") or factor.get("source_notes")),
                _display(target.get("confirmed_by_user") or factor.get("confirmed_by_user")),
            ])
    _style_table(sheet, header_row, widths=[34, 18, 18, 14, 20, 20, 14, 44, 12])
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL


def _build_f02(sheet, data: dict[str, Any]) -> None:
    f02 = data.get("f02") or {}
    enabled = bool(f02.get("comparison_workflow_enabled", True))
    sheet.append(["比較法調查估價資料", ""])
    sheet.merge_cells("A1:B1")
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append(["項目", "內容"])
    sheet.append(["比較法流程", "使用" if enabled else "本案未使用"])
    sheet.append(["比準地比較價格", _display(f02.get("benchmark_comparison_price"))])
    sheet.append(["比較價格決定說明", _display(f02.get("benchmark_notes"))])
    sheet.append(["備註", _display(f02.get("notes"))])
    sheet.append(["承辦人", _display(f02.get("handler_name"))])
    sheet.append(["課長", _display(f02.get("section_head_name"))])
    sheet.append(["主管", _display(f02.get("director_name"))])
    sheet.append(["估價人員", _display(f02.get("appraiser_name"))])
    sheet.append(["計算狀態", "已完成" if f02.get("calculation_status") == "CALCULATED" else "尚未完成"])
    sheet.append(["計算時間", _display(f02.get("calculated_at"))])
    sheet.append([])

    target_header = sheet.max_row + 1
    sheet.append([
        "比較標的順序",
        "正常單價",
        "交易日期",
        "日期調整率",
        "日期調整後單價",
        "區域因素調整率",
        "區域因素調整後單價",
        "個別因素調整率",
        "試算價格",
        "權重",
        "權重理由",
        "個別條件說明",
    ])
    targets = sorted(f02.get("comparison_targets") or [], key=lambda item: item.get("display_order") or 999)
    for target in targets:
        sheet.append([
            _display(target.get("display_order")),
            _display(target.get("normal_unit_price_snapshot")),
            _display(target.get("transaction_date_snapshot")),
            _display(target.get("time_adjustment_rate")),
            _display(target.get("date_adjusted_price")),
            _display(target.get("regional_adjustment_rate")),
            _display(target.get("regional_adjusted_price")),
            _display(target.get("individual_adjustment_rate")),
            _display(target.get("trial_price")),
            _display(target.get("weight")),
            _display(target.get("weight_reason")),
            _display(target.get("individual_condition_notes")),
        ])
    _style_table(sheet, target_header, widths=[14, 18, 14, 14, 18, 16, 20, 16, 18, 12, 36, 36])

    sheet.append([])
    factor_header = sheet.max_row + 1
    sheet.append([
        "比較標的順序",
        "個別因素",
        "比準地填報等級",
        "比準地確認等級",
        "比較標的填報等級",
        "比較標的確認等級",
        "調整率",
        "來源說明",
        "人工確認",
    ])
    for target in targets:
        for factor in target.get("individual_factors") or []:
            sheet.append([
                _display(target.get("display_order")),
                _factor_label(factor.get("factor_code"), individual=True),
                _display(factor.get("benchmark_reported_level")),
                _display(factor.get("benchmark_confirmed_level")),
                _display(factor.get("comparable_reported_level")),
                _display(factor.get("comparable_confirmed_level")),
                _display(factor.get("calculated_adjustment_rate")),
                _display(factor.get("source_notes")),
                _display(factor.get("confirmed_by_user")),
            ])
    _style_table(sheet, factor_header, widths=[14, 34, 18, 18, 20, 20, 14, 44, 12])
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
