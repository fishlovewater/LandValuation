from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from app.valuation.report_packages.factor_catalog import TEMPLATE_FACTORS


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEMPLATE_DIR = Path(__file__).with_name("templates") / "excel"
OFFICIAL_AUXILIARY_WORKBOOK = "table5_residential_regional_factors.xlsx"


class TemplateExportDefinition:
    def __init__(
        self,
        code: str,
        filename: str,
        sheet_name: str,
        title: str,
        *,
        output_filename: str | None = None,
    ) -> None:
        self.code = code
        self.filename = filename
        self.sheet_name = sheet_name
        self.title = title
        self.output_filename = output_filename or filename


TEMPLATE_EXPORTS = (
    TemplateExportDefinition("S01", "table3_section_survey.xlsx", "表3區段勘查表", "表3 地價區段勘查表"),
    TemplateExportDefinition("F02", "table4_comparison_method.xlsx", "表4比較法調查估價表", "表4 比較法調查估價表"),
    TemplateExportDefinition(
        "F02-RF",
        "table5_residential_regional_factors.xlsx",
        "表5-1區域因素明細表(住)",
        "表5-1 影響地價區域因素分析明細表（住宅用地）",
    ),
    TemplateExportDefinition(
        "F03",
        OFFICIAL_AUXILIARY_WORKBOOK,
        "114表14比準地地價估計表",
        "表14 比準地地價估計表",
        output_filename="table14_benchmark_land_value.xlsx",
    ),
    TemplateExportDefinition(
        "F04",
        OFFICIAL_AUXILIARY_WORKBOOK,
        "102表6宗地市價估計表格式",
        "表6 徵收土地宗地市價估計表",
        output_filename="table6_parcel_market_value.xlsx",
    ),
)

F04_FACTOR_ROWS = {
    "individual_area": 6,
    "individual_width": 7,
    "individual_depth": 8,
    "individual_shape": 9,
    "individual_street_frontage": 10,
    "individual_terrain": 11,
    "individual_road_type": 12,
    "individual_front_road_width": 13,
    "individual_school_proximity": 14,
    "individual_market_proximity": 15,
    "individual_park_proximity": 16,
    "individual_station_proximity": 17,
    "individual_commercial_district_proximity": 18,
    "individual_undesirable_facility": 19,
    "individual_parking_convenience": 20,
    "individual_land_use": 21,
    "individual_building_coverage_rate": 22,
    "individual_floor_area_ratio": 23,
    "individual_building_restriction": 24,
    "individual_other": 25,
}

S01_DIRECT_FIELDS = {
    "district_name", "district_boundary", "survey_date", "urban_plan_status",
    "land_use_zone", "building_coverage_rate", "floor_area_ratio",
    "prohibited_building", "restricted_building", "main_road_name",
    "main_road_width_m", "average_road_width_m", "notes", "site_opinion",
    "handler_name", "section_head_name", "director_name", "appraiser_name",
}
LOCATION_FIELD_ALIASES = {
    "administrative_area": "district_name",
    "zone_boundary_description": "district_boundary",
    "urban_plan_scope": "urban_plan_status",
    "land_use_zone_category": "land_use_zone",
    "building_coverage_ratio": "building_coverage_rate",
    "building_prohibition_status": "prohibited_building",
    "building_restriction_status": "restricted_building",
    "building_restriction_details": "restricted_building",
    "internal_road_width_m": "average_road_width_m",
    "average_internal_road_width_m": "average_road_width_m",
    "section_chief_name": "section_head_name",
}


def _display(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (dict, list, tuple)):
        return ", ".join(str(_display(item)) for item in value)
    return value


def _number(value: Any) -> Decimal | None:
    """Convert a known numeric export field without coercing identifiers."""

    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value
    try:
        return Decimal(str(value))
    except (ValueError, TypeError):
        return None


def _set(ws, cell: str, value: Any, rows: list[dict[str, Any]], field: str, note: str = "") -> None:
    rendered = _display(value)
    if rendered not in (None, "", []):
        target = ws[cell]
        if isinstance(target, MergedCell):
            for merged_range in ws.merged_cells.ranges:
                if target.coordinate in merged_range:
                    target = ws.cell(merged_range.min_row, merged_range.min_col)
                    break
        if not isinstance(target, MergedCell):
            target.value = rendered
    rows.append({"field_rules_field": field, "template_cell": cell, "value": rendered, "note": note})

def _append_audit_sheet(wb, rows: list[dict[str, Any]], title: str) -> None:
    # 正式輸出必須維持官方 workbook 結構，不再額外新增系統自製工作表。
    # 欄位來源與覆寫軌跡由既有結構化資料／快照保存，不寫回官方 Excel 版型。
    del wb, rows, title


def _target_items(f02: dict) -> list[dict]:
    return sorted(f02.get("comparison_targets") or [], key=lambda item: item.get("display_order", 0))[:3]


def _values(location: dict | None) -> dict[str, Any]:
    return dict((location or {}).get("values") or {})


def _label(location: dict | None) -> str | None:
    if location is None:
        return None
    return str(location.get("label") or location.get("address") or "").strip() or None


def _location_roles(locations: list[dict] | None) -> tuple[dict | None, list[dict]]:
    active = list(locations or [])
    benchmark = next((item for item in active if item.get("is_benchmark_location")), None)
    return benchmark, [item for item in active if item is not benchmark][:3]


def _location_value(location: dict | None, field: str, fallback: Any = None) -> Any:
    values = _values(location)
    return values.get(field) if values.get(field) not in (None, "", []) else fallback


def _append_location_audit(rows: list[dict[str, Any]], locations: list[dict] | None) -> None:
    for location in locations or []:
        role = "比準地" if location.get("is_benchmark_location") else "比較標的"
        label = _label(location) or "未命名地點"
        for field, value in _values(location).items():
            if value not in (None, "", []):
                rows.append({
                    "field_rules_field": field,
                    "template_cell": "地點彙整",
                    "value": _display(value),
                    "note": f"{role}：{label}",
                })


def _location_s01(s01: dict, location: dict | None) -> dict:
    if location is None:
        return s01
    data = dict(s01)
    values = _values(location)
    for source, target in LOCATION_FIELD_ALIASES.items():
        if values.get(source) not in (None, "", []):
            values[target] = values[source]
    for field in S01_DIRECT_FIELDS:
        if values.get(field) not in (None, "", []):
            data[field] = values[field]
    observations = {str(item.get("item_code")): dict(item) for item in (data.get("observations") or [])}
    for factor in TEMPLATE_FACTORS:
        value = values.get(factor.code)
        if value not in (None, "", []):
            observations[factor.code] = {"item_code": factor.code, "raw_value": value, "source_notes": "地點 AI／人工確認值"}
    data["observations"] = list(observations.values())
    return data


def _build_s01(wb, case: dict, s01: dict, location: dict | None = None) -> None:
    ws = wb["表3區段勘查表"]
    rows: list[dict[str, Any]] = []
    s01 = _location_s01(s01, location)
    location_note = _label(location) or "全案"
    _set(ws, "B3", case.get("valuation_base_date"), rows, "valuation_base_date")
    _set(ws, "G3", s01.get("price_zone_no") or case.get("price_zone_no"), rows, "price_zone_no")
    _set(ws, "I3", s01.get("district_boundary"), rows, "zone_boundary_description")
    _set(ws, "H4", s01.get("urban_plan_status"), rows, "urban_plan_scope")
    _set(ws, "H5", s01.get("land_use_zone"), rows, "land_use_zone_category")
    _set(ws, "H6", s01.get("building_coverage_rate"), rows, "building_coverage_ratio", "百分比")
    _set(ws, "H7", s01.get("floor_area_ratio"), rows, "floor_area_ratio", "百分比")
    _set(ws, "H8", s01.get("prohibited_building"), rows, "prohibited_building")
    _set(ws, "H9", s01.get("restricted_building"), rows, "restricted_building")
    _set(ws, "G11", s01.get("main_road_name"), rows, "main_road_name")
    _set(ws, "J11", s01.get("main_road_width_m"), rows, "main_road_width_m", "公尺")
    _set(ws, "I12", s01.get("average_road_width_m"), rows, "average_internal_road_width_m", "公尺")
    _set(ws, "B45", s01.get("survey_date"), rows, "survey_date")
    _set(ws, "I45", s01.get("handler_name"), rows, "handler_name")
    _set(ws, "N45", s01.get("section_head_name"), rows, "section_head_name")
    _set(ws, "S45", s01.get("director_name"), rows, "director_name")
    _set(ws, "S46", s01.get("appraiser_name"), rows, "appraiser_name")
    observation_cells = {
        "mass_transit_proximity": "H13", "station_proximity": "H14", "interchange_proximity": "I19", "road_plan": "D23",
        "drainage": "D27", "terrain": "D28", "market_proximity": "I39", "park_proximity": "I42",
        "tourist_facility_proximity": "R36", "parking_convenience": "O6", "power_gas_facility": "P14",
        "funeral_facility": "Q18", "waste_facility": "Q22", "environmental_pollution": "Q25",
        "department_store": "Q30", "financial_institution": "Q32", "entertainment_facility": "Q34",
        "exhibition_hotel": "Q36", "pedestrian_flow": "Q38", "vacancy_rate": "Q39", "other": "L40",
    }
    for item in s01.get("observations") or []:
        code = str(item.get("item_code") or "")
        cell = observation_cells.get(code)
        if cell:
            value = item.get("raw_value")
            if item.get("facility_name") and item.get("walking_distance_m") is not None:
                value = f"{item['facility_name']}（{item['walking_distance_m']}m）"
            _set(ws, cell, value, rows, code, item.get("source_notes") or "已確認觀察")
        else:
            rows.append({"field_rules_field": code, "template_cell": "AI欄位對應", "value": _display(item.get("raw_value")), "note": "範本沒有一對一空格，保留於對應頁。"})
    _set(ws, "A47", s01.get("notes"), rows, "case_note")
    _append_location_audit(rows, [location] if location else None)
    _append_audit_sheet(wb, rows, f"S01／表3 field_rules 對應（{location_note}）")


def _factor_value(location: dict | None, code: str, fallback: Any = None) -> Any:
    return _location_value(location, code, fallback)


def _build_f02(wb, case: dict, f02: dict, f02rf: dict, locations: list[dict] | None = None) -> None:
    ws = wb["表4比較法調查估價表"]
    rows: list[dict[str, Any]] = []
    benchmark_location, location_targets = _location_roles(locations)
    fallback_targets = _target_items(f02)
    _set(ws, "L1", case.get("valuation_base_date"), rows, "valuation_base_date")
    _set(ws, "P1", case.get("case_no"), rows, "case_no")
    _set(ws, "D4", _label(benchmark_location) or f02.get("benchmark_land_id"), rows, "benchmark_land_id", "前段選擇的比準地")
    for index in range(3):
        target = location_targets[index] if index < len(location_targets) else None
        fallback = fallback_targets[index] if index < len(fallback_targets) else {}
        base = ("G", "K", "O")[index]
        _set(ws, f"{base}4", _label(target) or fallback.get("transaction_no") or fallback.get("comparison_target_id"), rows, "instance_no", f"比較標的 {index + 1}")
    _set(ws, "D5", _location_value(benchmark_location, "normal_land_unit_price"), rows, "normal_land_unit_price", "比準地")
    _set(ws, "D6", _location_value(benchmark_location, "transaction_date"), rows, "transaction_date", "比準地")
    _set(ws, "D8", f02rf.get("regional_adjustment_rates"), rows, "regional_factor_rate")
    factor_cells = ["D", "G", "K", "O"]
    individual_row = {
        "individual_area": 9, "individual_width": 10, "individual_depth": 11, "individual_shape": 12,
        "individual_street_frontage": 13, "individual_terrain": 14, "individual_road_type": 15,
        "individual_front_road_width": 16, "individual_school_proximity": 17, "individual_market_proximity": 18,
        "individual_park_proximity": 19, "individual_station_proximity": 20,
        "individual_commercial_district_proximity": 21, "individual_undesirable_facility": 22,
        "individual_parking_convenience": 23, "individual_land_use": 24,
        "individual_building_coverage_rate": 25, "individual_floor_area_ratio": 26,
        "individual_building_restriction": 27, "individual_other": 28,
    }
    for code, row in individual_row.items():
        _set(ws, f"D{row}", _factor_value(benchmark_location, code), rows, code, "比準地")
        for index in range(3):
            target = location_targets[index] if index < len(location_targets) else None
            fallback = fallback_targets[index] if index < len(fallback_targets) else {}
            legacy = next((item.get("comparable_confirmed_level") or item.get("comparable_reported_level") for item in fallback.get("individual_factors") or [] if item.get("factor_code") == code), None)
            _set(ws, f"{factor_cells[index + 1]}{row}", _factor_value(target, code, legacy), rows, code, f"比較標的 {index + 1}")
    for index in range(3):
        target = location_targets[index] if index < len(location_targets) else None
        fallback = fallback_targets[index] if index < len(fallback_targets) else {}
        col = factor_cells[index + 1]
        _set(ws, f"{col}29", _location_value(target, "absolute_adjustment_total", fallback.get("total_adjustment_absolute")), rows, "absolute_adjustment_total")
        _set(ws, f"{col}30", _location_value(target, "regional_adjusted_price", fallback.get("regional_adjusted_price")), rows, "regional_adjusted_price")
        _set(ws, f"{col}31", _location_value(target, "trial_price", fallback.get("trial_price")), rows, "trial_price")
        _set(ws, f"{col}32", _location_value(target, "comparison_weight", fallback.get("weight")), rows, "comparison_weight")
    _set(ws, "D32", f02.get("benchmark_comparison_price"), rows, "benchmark_comparison_price")
    _set(ws, "B34", f02.get("notes"), rows, "case_note")
    _append_location_audit(rows, locations)
    _append_audit_sheet(wb, rows, "F02／表4 field_rules 對應（全部地點）")


def _factor_target(factor: dict, fallback: dict, display_order: int) -> dict:
    fallback_id = fallback.get("comparison_target_id")
    return next(
        (
            item
            for item in factor.get("targets") or []
            if item.get("display_order") == display_order
            or (
                fallback_id is not None
                and str(item.get("comparison_target_id")) == str(fallback_id)
            )
        ),
        {},
    )


def _build_f02_rf(wb, case: dict, f02rf: dict, f02: dict, locations: list[dict] | None = None) -> None:
    if str(case.get("land_use_type") or "").upper() != "RESIDENTIAL":
        raise FileNotFoundError(
            "目前只提供官方表5-1住宅用地範本；非住宅用地不得由住宅範本改造產生。"
        )
    ws = wb["表5-1區域因素明細表(住)"]
    rows: list[dict[str, Any]] = []
    benchmark_location, location_targets = _location_roles(locations)
    fallback_targets = _target_items(f02)
    _set(ws, "B2", case.get("case_no"), rows, "case_no")
    _set(ws, "C3", _label(benchmark_location) or f02rf.get("benchmark_land_id"), rows, "benchmark_land_id", "前段選擇的比準地")
    for index in range(3):
        target = location_targets[index] if index < len(location_targets) else None
        fallback = fallback_targets[index] if index < len(fallback_targets) else {}
        _set(ws, ("E3", "H3", "K3")[index], _label(target) or fallback.get("transaction_no") or fallback.get("comparison_target_id"), rows, "instance_no", f"比較標的 {index + 1}")
    row_map = {
        "urban_plan_status": 5, "land_use_zone": 6, "building_coverage_rate": 7, "floor_area_ratio": 8,
        "prohibited_building": 9, "restricted_building": 10, "main_road_width": 12, "average_road_width": 13,
        "mass_transit_proximity": 14, "station_proximity": 15, "interchange_proximity": 16, "road_plan": 17,
        "drainage": 19, "terrain": 20, "market_proximity": 27, "park_proximity": 28,
        "tourist_facility_proximity": 29, "parking_convenience": 30, "power_gas_facility": 34,
        "funeral_facility": 35, "waste_facility": 36, "environmental_pollution": 38, "department_store": 40,
    }
    factor_rows = {str(item.get("factor_code")): item for item in (f02rf.get("factor_rows") or [])}
    target_cols = ("E", "H", "K")
    rate_cols = ("G", "J", "M")
    for code, row in row_map.items():
        factor = factor_rows.get(code, {})
        benchmark_level = _factor_value(
            benchmark_location,
            code,
            factor.get("benchmark_confirmed_level") or factor.get("benchmark_reported_level"),
        )
        _set(ws, f"C{row}", benchmark_level, rows, code, "比準地")
        for index in range(3):
            target = location_targets[index] if index < len(location_targets) else None
            fallback = fallback_targets[index] if index < len(fallback_targets) else {}
            target_data = _factor_target(factor, fallback, index + 1)
            location_values = _values(target)
            calculated_levels = location_values.get("regional_factor_levels") or {}
            calculated_rates = location_values.get("regional_factor_adjustment_rates") or {}
            target_level = _factor_value(
                target,
                code,
                calculated_levels.get(code) or target_data.get("confirmed_level") or target_data.get("reported_level"),
            )
            _set(ws, f"{target_cols[index]}{row}", target_level, rows, code, f"比較標的 {index + 1}")
            rate_val = calculated_rates.get(code, target_data.get("calculated_adjustment_rate"))
            rate = _number(rate_val)
            _set(ws, f"{rate_cols[index]}{row}", rate, rows, f"{code}.adjustment_rate", "正式計算結果")
    _set(ws, "C40", f02rf.get("other_influences"), rows, "other")
    _set(ws, "C44", f02rf.get("notes"), rows, "case_note")
    _append_location_audit(rows, locations)
    _append_audit_sheet(wb, rows, "F02-RF／表5 field_rules 對應（全部地點）")


def _clear_value(ws, cell: str) -> None:
    target = ws[cell]
    if isinstance(target, MergedCell):
        for merged_range in ws.merged_cells.ranges:
            if target.coordinate in merged_range:
                target = ws.cell(merged_range.min_row, merged_range.min_col)
                break
    if not isinstance(target, MergedCell):
        target.value = None


def _build_f03(wb, case: dict, f03: dict) -> None:
    ws = wb["114表14比準地地價估計表"]
    rows: list[dict[str, Any]] = []
    for cell in ("B2", "I2", "A10", "J11"):
        _clear_value(ws, cell)
    for row in range(5, 10):
        for col in "ABCDEFGHIJK":
            _clear_value(ws, f"{col}{row}")

    _set(ws, "B2", case.get("case_no"), rows, "case_no")
    _set(ws, "I2", case.get("valuation_base_date"), rows, "valuation_base_date")
    for index, item in enumerate((f03.get("rows") or [])[:5], start=5):
        _set(ws, f"A{index}", item.get("price_zone_no"), rows, "price_zone_no")
        _set(ws, f"B{index}", item.get("benchmark_land_no"), rows, "benchmark_land_no")
        _set(ws, f"C{index}", item.get("district_name"), rows, "district_name")
        _set(ws, f"D{index}", item.get("section_subsection_name"), rows, "section_subsection_name")
        _set(ws, f"E{index}", item.get("land_no"), rows, "land_no")
        _set(ws, f"F{index}", _number(item.get("comparison_price")), rows, "comparison_price")
        _set(ws, f"G{index}", _number(item.get("comparison_weight")), rows, "comparison_weight")
        _set(ws, f"H{index}", _number(item.get("income_price")) if item.get("income_price") is not None else "-", rows, "income_price")
        _set(ws, f"I{index}", _number(item.get("income_weight")), rows, "income_weight")
        _set(ws, f"J{index}", _number(item.get("benchmark_land_price")), rows, "benchmark_land_price")
        _set(ws, f"K{index}", item.get("decision_reason"), rows, "decision_reason")

    footer = f03.get("footer") or {}
    footer_text = (
        f"填寫日期：{_display(footer.get('filled_date')) or ''}    "
        f"承辦員：{footer.get('handler_name') or ''}    "
        f"課（股）長：{footer.get('section_head_name') or ''}    "
        f"主任（局、處長）：{footer.get('director_name') or ''}"
    )
    _set(ws, "A10", footer_text, rows, "signatures")
    _set(ws, "J11", footer.get("appraiser_name"), rows, "appraiser_name")
    _append_audit_sheet(wb, rows, "F03／表14 field_rules 對應")


def _build_f04(wb, case: dict, f04: dict) -> None:
    ws = wb["102表6宗地市價估計表格式"]
    rows: list[dict[str, Any]] = []
    condition_cols = ("F", "J", "N", "R", "V")
    rate_cols = ("I", "M", "Q", "U", "Y")
    serial_cells = ("I2", "M2", "Q2", "U2", "Y2")

    for cell in ("O1", "W1", "Y1", "E2", "C4", "C5", "C30", "A31", "M32"):
        _clear_value(ws, cell)
    for row in F04_FACTOR_ROWS.values():
        _clear_value(ws, f"C{row}")
        for condition_col, rate_col in zip(condition_cols, rate_cols):
            _clear_value(ws, f"{condition_col}{row}")
            _clear_value(ws, f"{rate_col}{row}")
    for index in range(5):
        for row in (26, 27, 28, 29):
            _clear_value(ws, f"{condition_cols[index]}{row}")
        _clear_value(ws, serial_cells[index])
        _clear_value(ws, f"{condition_cols[index]}4")

    _set(ws, "O1", case.get("case_no"), rows, "case_no")
    benchmark = f04.get("benchmark") or {}
    _set(ws, "E2", benchmark.get("benchmark_land_no"), rows, "benchmark_land_no")
    _set(ws, "C4", benchmark.get("parcel_display"), rows, "benchmark_parcel")
    _set(ws, "C5", _number(f04.get("benchmark_land_price")), rows, "benchmark_land_price")
    benchmark_factors = {
        str(item.get("factor_code")): item for item in f04.get("benchmark_factor_rows") or []
    }
    for code, row in F04_FACTOR_ROWS.items():
        _set(ws, f"C{row}", (benchmark_factors.get(code) or {}).get("condition_value"), rows, f"benchmark.{code}")

    snapshot_rows = {
        str(item.get("parcel_id")): item
        for item in ((f04.get("calculation_snapshot") or {}).get("rows") or [])
    }
    parcels = (f04.get("parcel_rows") or [])[:5]
    if parcels:
        _set(ws, "W1", parcels[0].get("parcel_serial"), rows, "parcel_serial_start")
        _set(ws, "Y1", parcels[-1].get("parcel_serial"), rows, "parcel_serial_end")
    for index, item in enumerate(parcels):
        condition_col = condition_cols[index]
        rate_col = rate_cols[index]
        _set(ws, serial_cells[index], item.get("parcel_serial"), rows, "parcel_serial", f"宗地 {index + 1}")
        _set(ws, f"{condition_col}4", item.get("parcel_display"), rows, "parcel_display", f"宗地 {index + 1}")
        factor_rows = {str(row.get("factor_code")): row for row in item.get("factor_rows") or []}
        for code, row_no in F04_FACTOR_ROWS.items():
            factor = factor_rows.get(code) or {}
            _set(ws, f"{condition_col}{row_no}", factor.get("condition_value"), rows, f"{code}.condition", f"宗地 {index + 1}")
            _set(ws, f"{rate_col}{row_no}", _number(factor.get("difference_rate")), rows, f"{code}.difference_rate", f"宗地 {index + 1}")
        snapshot = snapshot_rows.get(str(item.get("parcel_id"))) or {}
        _set(ws, f"{condition_col}26", _number(item.get("parcel_adjustment_rate")), rows, "parcel_adjustment_rate", f"宗地 {index + 1}")
        _set(ws, f"{condition_col}27", _number(snapshot.get("trial_unit_price")), rows, "trial_unit_price", f"宗地 {index + 1}")
        _set(ws, f"{condition_col}28", _number(item.get("parcel_unit_price")), rows, "parcel_unit_price", f"宗地 {index + 1}")
        _set(ws, f"{condition_col}29", item.get("notes"), rows, "parcel_note", f"宗地 {index + 1}")

    _set(ws, "C30", f04.get("notes"), rows, "case_note")
    footer = f04.get("footer") or {}
    footer_text = (
        f"填寫日期：{_display(footer.get('filled_date')) or ''}    "
        f"承辦員：{footer.get('handler_name') or ''}    "
        f"課（股）長：{footer.get('section_head_name') or ''}    "
        f"主任（局、處長）：{footer.get('director_name') or ''}"
    )
    _set(ws, "A31", footer_text, rows, "signatures")
    _set(ws, "M32", footer.get("appraiser_name"), rows, "appraiser_name")
    _append_audit_sheet(wb, rows, "F04／表6 field_rules 對應")


def build_template_export_xlsx(*, code: str, case: dict, pages: dict[str, dict], locations: list[dict] | None = None, location: dict | None = None) -> bytes:
    definition = next(item for item in TEMPLATE_EXPORTS if item.code == code)
    land_use = str(case.get("land_use_type") or "").upper()
    if code == "F02-RF" and land_use != "RESIDENTIAL":
        raise FileNotFoundError(
            f"缺少 {land_use or 'UNKNOWN'} 用地的官方表5範本；"
            "系統不會再由住宅表5-1自行改造成其他用地範本。"
        )
    path = TEMPLATE_DIR / definition.filename
    if not path.exists():
        raise FileNotFoundError(f"找不到 Excel 範本：{path}")
    wb = load_workbook(path)
    if definition.sheet_name not in wb.sheetnames:
        raise FileNotFoundError(
            f"官方 Excel 範本缺少工作表：{definition.sheet_name}（{path.name}）"
        )
    if code == "S01":
        _build_s01(wb, case, pages["S01"], location)
    elif code == "F02":
        _build_f02(wb, case, pages["F02"], pages["F02-RF"], locations)
    elif code == "F02-RF":
        _build_f02_rf(wb, case, pages["F02-RF"], pages["F02"], locations)
    elif code == "F03":
        _build_f03(wb, case, pages["F03"])
    elif code == "F04":
        _build_f04(wb, case, pages["F04"])
    else:
        raise ValueError(f"不支援的正式 Excel 表單：{code}")
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
