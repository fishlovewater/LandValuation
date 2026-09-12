from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from typing import Iterable, Mapping

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.valuation.extraction.field_analysis import FIELD_ANALYSIS_FIELDS
from app.valuation.official_field_catalog import OFFICIAL_PDF_FIELDS
from app.valuation.report_packages.factor_catalog import INDIVIDUAL_FACTORS, TEMPLATE_FACTORS

EXPLICIT_FIELD_LABELS: dict[str, str] = {
    # 一、宗地本身物理條件與個別因素 (F02 表4 比較法)
    "individual_area": "面積 (M²)",
    "individual_width": "寬度 (M)",
    "individual_depth": "深度 (M)",
    "individual_shape": "形狀",
    "individual_street_frontage": "臨街情形",
    "individual_terrain": "地勢",

    # 二、道路與交通條件（個別因素）
    "individual_road_type": "道路種類",
    "individual_front_road_width": "面前道路寬度 (M)",

    # 三、土地利用現況與改良
    "building_site_improvement_type": "建築基地改良項目",
    "building_site_improvements": "建築基地改良明細",
    "farmland_improvement_type": "農地改良項目",
    "farmland_improvements": "農地改良明細",
    "building_status": "房屋建築現況",
    "dominant_building_type": "建築形態",
    "current_land_use_type": "主要土地利用現況",

    # 四、特定周邊環境與工商活動
    "individual_school_proximity": "接近學校之程度",
    "school_proximity": "接近學校之程度",
    "individual_commercial_district_proximity": "接近商圈之程度",
    "commercial_district_proximity": "接近商圈之程度",
    "service_facility_proximity": "接近服務性設施之程度",
    "customer_traffic_level": "顧客通行量之多寡",
    "pedestrian_flow": "顧客通行量之多寡",
    "shop_contiguity_level": "店舖之毗連狀態",
    "vacancy_rate": "店舖之毗連狀態",

    # 五、嫌惡設施（獨立細項）
    "individual_undesirable_facility": "嫌惡設施之有無及接近程度",
    "substation_high_voltage_tower": "變電所或高壓鐵塔",
    "gas_oil_tank": "瓦斯槽或儲油槽",
    "cemetery": "墓地",
    "funeral_home": "殯儀館",
    "crematorium": "火葬場",
    "columbarium": "納骨塔",
    "funeral_facility": "殯葬設施之有無及接近程度",
    "sewage_plant": "污水處理場",
    "landfill_incinerator": "垃圾場、掩埋場或焚化爐",
    "waste_facility": "廢棄物處理設施之有無及接近程度",

    # 六、環境污染（獨立細項）
    "water_pollution": "水污染",
    "noise_pollution": "噪音污染",
    "air_pollution": "廢氣污染",
    "waste_pollution": "廢棄物污染",
    "environmental_pollution": "水污染、噪音污染、廢氣污染、廢棄物污染等之有無及接近程度",

    # 七、估價計算與小計欄位
    "category_subtotal": "百分比小計",
    "subtotal_percentage": "百分比小計",
    "total_regional_factor_correction": "影響地價區域因素總修正數",
    "total_correction_percentage": "區域因素總修正數",
    "individual_factor_total": "個別因素調整百分率合計",
    "absolute_adjustment_total": "調整百分率絕對值加總",
    "trial_price": "試算價格（元／㎡）",
    "comparison_weight": "比較標的權重",
    "benchmark_comparison_price": "比準地比較價格（元／㎡）",

    # 八、名稱修復（修正括號截斷）
    "drainage_level": "保（排）水之良否",
}


def _text(value: object | None) -> str:
    if value is None:
        return ""
    return str(value)


def get_field_label_zh(form_code: str, field_name: str) -> str:
    """Return the official Chinese label for a given form code and field name."""
    if field_name in EXPLICIT_FIELD_LABELS:
        return EXPLICIT_FIELD_LABELS[field_name]

    pdf_fields = OFFICIAL_PDF_FIELDS.get(form_code, ())
    for field in pdf_fields:
        if field.code == field_name:
            return field.label

    for factor in INDIVIDUAL_FACTORS:
        if factor.code == field_name:
            return factor.label

    for factor in TEMPLATE_FACTORS:
        if factor.code == field_name:
            return factor.label

    catalog = FIELD_ANALYSIS_FIELDS.get(form_code, {})
    raw_desc = catalog.get(field_name)
    if raw_desc:
        clean = re.split(r"[;；\n]", raw_desc)[0].strip()
        if clean:
            return clean

    for cat in FIELD_ANALYSIS_FIELDS.values():
        if field_name in cat:
            clean = re.split(r"[;；\n]", cat[field_name])[0].strip()
            if clean:
                return clean

    return field_name


STANDARD_FORM_ORDER = ("F03", "F01", "S01", "F02-RF", "F02", "F04")


# These are deliberately limited to formulas that already exist in the
# valuation workflow.  A generated value is review-only: it never overwrites a
# document-grounded candidate or writes back to a formal form.
CALCULATED_FIELDS: dict[str, tuple[str, ...]] = {
    "S01": ("average_internal_road_width_m",),
    "F01": ("normal_land_total_price", "normal_land_unit_price"),
    "F02-RF": ("regional_adjustment_rate", "total_adjustment_rate"),
    "F02": (
        "adjusted_unit_price_raw",
        "adjusted_unit_price_display",
        "absolute_adjustment_total",
        "trial_price_raw",
        "trial_price",
        "benchmark_comparison_price",
    ),
    "F03": ("weight_total", "weighted_value_raw", "benchmark_land_price"),
    "F04": ("parcel_unit_price", "parcel_total_value"),
}


@dataclass(frozen=True)
class FormulaFallback:
    status: str
    value: str | None
    formula: str
    note: str


def _resolved_candidate_value(candidates: Iterable[object]) -> object | None:
    """Return the most recent human-accepted, document-grounded value."""
    for candidate in reversed(list(candidates)):
        if getattr(candidate, "field_status", None) not in {"CONFIRMED", "APPLIED"}:
            continue
        value = getattr(candidate, "confirmed_value", None)
        if value is None:
            value = getattr(candidate, "extracted_value", None)
        if value is not None and str(value).strip() != "":
            return value
    return None


def _decimal(value: object | None) -> Decimal | None:
    if value is None:
        return None
    raw = str(value).strip().replace(",", "").replace("，", "")
    if not raw:
        return None
    is_percent = raw.endswith("%")
    if is_percent:
        raw = raw[:-1].strip()
    try:
        result = Decimal(raw)
    except (InvalidOperation, ValueError):
        return None
    return result / Decimal("100") if is_percent else result


def _format_decimal(value: Decimal, quantum: Decimal) -> str:
    return format(value.quantize(quantum, rounding=ROUND_HALF_UP), "f")


def _missing_note(form_code: str, names: Iterable[str], formula: str) -> FormulaFallback:
    labels = "、".join(get_field_label_zh(form_code, name) for name in names)
    return FormulaFallback(
        status="FORMULA_INPUT_MISSING",
        value=None,
        formula=formula,
        note=f"無法計算；缺少或格式無法辨識：{labels}。",
    )


def _formula_fallbacks(
    form_code: str,
    grouped: dict[tuple[str, str], list[object]],
) -> dict[str, FormulaFallback]:
    """Build only the safe formula fallbacks for fields with no direct answer."""
    direct = {
        field_name: _resolved_candidate_value(field_candidates)
        for (candidate_form, field_name), field_candidates in grouped.items()
        if candidate_form == form_code
    }
    values: dict[str, Decimal] = {
        name: value
        for name, raw in direct.items()
        if (value := _decimal(raw)) is not None
    }
    results: dict[str, FormulaFallback] = {}

    def has_direct(field_name: str) -> bool:
        return direct.get(field_name) is not None

    def missing(*names: str) -> list[str]:
        return [name for name in names if name not in values]

    def add(
        field_name: str,
        value: Decimal | None,
        formula: str,
        *,
        quantum: Decimal = Decimal("0.01"),
        missing_names: Iterable[str] = (),
    ) -> None:
        if has_direct(field_name):
            return
        missing_list = list(missing_names)
        if missing_list:
            results[field_name] = _missing_note(form_code, missing_list, formula)
            return
        if value is None:
            return
        values[field_name] = value
        results[field_name] = FormulaFallback(
            status="FORMULA_CALCULATED",
            value=_format_decimal(value, quantum),
            formula=formula,
            note="未找到已確認的直接答案，依既有公式以已確認輸入值計算。",
        )

    if form_code == "F01":
        names = (
            "transaction_total_price",
            "building_price_deduction",
            "special_transaction_adjustment",
            "land_area_sqm",
        )
        missing_names = missing(*names)
        normal_total = None
        if not missing_names:
            normal_total = (values[names[0]] - values[names[1]]) * (
                Decimal("1") + values[names[2]]
            )
        add(
            "normal_land_total_price",
            normal_total,
            "(transaction_total_price - building_price_deduction) × (1 + special_transaction_adjustment)",
            missing_names=missing_names,
        )
        unit_missing = missing("normal_land_total_price", "land_area_sqm")
        unit_price = None if unit_missing else values["normal_land_total_price"] / values["land_area_sqm"]
        add(
            "normal_land_unit_price",
            unit_price,
            "normal_land_total_price ÷ land_area_sqm",
            missing_names=unit_missing,
        )

    elif form_code == "F02":
        date_names = ("normal_land_unit_price", "date_adjustment_rate")
        date_missing = missing(*date_names)
        date_adjusted = None
        if not date_missing:
            date_adjusted = values[date_names[0]] * (Decimal("1") + values[date_names[1]])
        add(
            "adjusted_unit_price_raw",
            date_adjusted,
            "normal_land_unit_price × (1 + date_adjustment_rate)",
            missing_names=date_missing,
        )
        display_missing = missing("adjusted_unit_price_raw")
        add(
            "adjusted_unit_price_display",
            None if display_missing else values["adjusted_unit_price_raw"],
            "round(adjusted_unit_price_raw, 0)",
            quantum=Decimal("1"),
            missing_names=display_missing,
        )

        trial_names = (
            "normal_land_unit_price",
            "date_adjustment_rate",
            "regional_factor_rate",
            "individual_factor_total",
        )
        trial_missing = missing(*trial_names)
        trial = None
        if not trial_missing:
            trial = (
                values["normal_land_unit_price"]
                * (Decimal("1") + values["date_adjustment_rate"])
                * (Decimal("1") + values["regional_factor_rate"])
                * (Decimal("1") + values["individual_factor_total"])
            )
        add(
            "trial_price_raw",
            trial,
            "normal_land_unit_price × (1 + date_adjustment_rate) × (1 + regional_factor_rate) × (1 + individual_factor_total)",
            missing_names=trial_missing,
        )
        trial_display_missing = missing("trial_price_raw")
        add(
            "trial_price",
            None if trial_display_missing else values["trial_price_raw"],
            "round(trial_price_raw, 0)",
            quantum=Decimal("1"),
            missing_names=trial_display_missing,
        )
        absolute_names = (
            "date_adjustment_rate",
            "regional_factor_rate",
            "individual_factor_total",
        )
        absolute_missing = missing(*absolute_names)
        absolute = None
        if not absolute_missing:
            absolute = sum(abs(values[name]) for name in absolute_names)
        add(
            "absolute_adjustment_total",
            absolute,
            "|date_adjustment_rate| + |regional_factor_rate| + |individual_factor_total|",
            quantum=Decimal("0.000001"),
            missing_names=absolute_missing,
        )
        benchmark_missing = missing("trial_price", "comparison_weight")
        benchmark = None
        if not benchmark_missing and values["comparison_weight"] == Decimal("1"):
            benchmark = values["trial_price"]
        elif not benchmark_missing:
            benchmark_missing = ["comparison_weight（所有比較標的權重合計 100%）"]
        add(
            "benchmark_comparison_price",
            benchmark,
            "Σ(trial_price × comparison_weight)，並依規則取整數",
            quantum=Decimal("1"),
            missing_names=benchmark_missing,
        )

    elif form_code == "F03":
        weight_names = ("comparison_weight", "income_weight")
        total_missing = missing(*weight_names)
        weight_total = None if total_missing else values["comparison_weight"] + values["income_weight"]
        add(
            "weight_total",
            weight_total,
            "comparison_weight + income_weight",
            quantum=Decimal("0.000001"),
            missing_names=total_missing,
        )
        price_names: list[str] = ["comparison_weight", "income_weight"]
        if values.get("comparison_weight", Decimal("0")) > 0:
            price_names.append("comparison_price")
        if values.get("income_weight", Decimal("0")) > 0:
            price_names.append("income_price")
        price_missing = missing(*price_names)
        raw = None
        if not price_missing and values.get("weight_total") == Decimal("1"):
            raw = (
                values.get("comparison_price", Decimal("0")) * values["comparison_weight"]
                + values.get("income_price", Decimal("0")) * values["income_weight"]
            )
        elif not price_missing:
            price_missing = ["weight_total（比較法與收益法權重合計 100%）"]
        add(
            "weighted_value_raw",
            raw,
            "comparison_price × comparison_weight + income_price × income_weight",
            missing_names=price_missing,
        )
        benchmark_missing = missing("weighted_value_raw")
        add(
            "benchmark_land_price",
            None if benchmark_missing else values["weighted_value_raw"],
            "round(weighted_value_raw, 2)",
            missing_names=benchmark_missing,
        )

    elif form_code == "F04":
        unit_names = ("benchmark_land_price", "parcel_adjustment_rate")
        unit_missing = missing(*unit_names)
        unit_price = None
        if not unit_missing:
            unit_price = values["benchmark_land_price"] * (
                Decimal("1") + values["parcel_adjustment_rate"]
            )
        add(
            "parcel_unit_price",
            unit_price,
            "benchmark_land_price × (1 + parcel_adjustment_rate)",
            missing_names=unit_missing,
        )
        total_names = ("parcel_unit_price", "parcel_area_sqm", "ownership_ratio")
        total_missing = missing(*total_names)
        total_value = None
        if not total_missing:
            total_value = (
                values["parcel_unit_price"]
                * values["parcel_area_sqm"]
                * values["ownership_ratio"]
            )
        add(
            "parcel_total_value",
            total_value,
            "parcel_unit_price × parcel_area_sqm × ownership_ratio",
            quantum=Decimal("1"),
            missing_names=total_missing,
        )

    elif form_code == "F02-RF":
        for field_name in CALCULATED_FIELDS[form_code]:
            if not has_direct(field_name):
                results[field_name] = _missing_note(
                    form_code,
                    ("rule_version_id", "confirmed_factor_levels", "comparison_targets"),
                    "依已發布規則版本的因素級距計算",
                )

    elif form_code == "S01" and not has_direct("average_internal_road_width_m"):
        results["average_internal_road_width_m"] = FormulaFallback(
            status="FORMULA_INPUT_MISSING",
            value=None,
            formula="無可驗證的既有公式",
            note="無法計算；需由來源文件或人工確認區段內道路平均寬度。",
        )

    return results


def build_confirmation_export_xlsx(
    *,
    case_no: str,
    case_title: str,
    generated_at: datetime,
    candidates: Iterable[object],
    manual_values: Mapping[str, Mapping[str, object]] | None = None,
) -> bytes:
    """Build a review export with direct evidence before formula fallbacks."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "欄位確認結果"

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="B7C9D6"),
        right=Side(style="thin", color="B7C9D6"),
        top=Side(style="thin", color="B7C9D6"),
        bottom=Side(style="thin", color="B7C9D6"),
    )
    align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)

    ws.append(["欄位確認結果清單"])
    ws.cell(row=1, column=1).font = Font(size=14, bold=True)
    ws.append([f"案件：{case_no} {case_title}"])
    ws.append([f"產生時間（UTC）：{generated_at.isoformat()} 此文件包含完整表單欄位對照，未辨識到之欄位自動留白。"])
    ws.append([])

    headers = [
        "表單",
        "欄位代碼",
        "候選群組",
        "不同候選答案數",
        "欄位中文名稱",
        "狀態",
        "確認/修正值",
        "擷取原始值",
        "來源證據原文",
        "頁碼",
        "分析來源",
        "信心度",
        "建立時間",
        "計算／缺值備註",
    ]
    header_row_idx = 5
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    candidates_list = list(candidates)
    manual_values = manual_values or {}
    grouped: dict[tuple[str, str], list[object]] = defaultdict(list)
    for c in candidates_list:
        code = getattr(c, "form_code", "")
        name = getattr(c, "field_name", "")
        if code and name:
            grouped[(code, name)].append(c)

    candidate_form_codes = {getattr(c, "form_code", "") for c in candidates_list if getattr(c, "form_code", "")}
    ordered_forms = list(STANDARD_FORM_ORDER)
    for f in sorted(candidate_form_codes):
        if f not in ordered_forms:
            ordered_forms.append(f)

    for form_code in ordered_forms:
        defined_catalog = FIELD_ANALYSIS_FIELDS.get(form_code, {})
        defined_fields = list(defined_catalog.keys())
        if not defined_fields and form_code in OFFICIAL_PDF_FIELDS:
            defined_fields = [f.code for f in OFFICIAL_PDF_FIELDS[form_code]]

        candidate_field_names = [name for (fc, name) in grouped if fc == form_code]
        all_fields = list(defined_fields)
        for name in manual_values.get(form_code, {}):
            if name not in all_fields:
                all_fields.append(name)
        for field_name in CALCULATED_FIELDS.get(form_code, ()):
            if field_name not in all_fields:
                all_fields.append(field_name)
        for name in candidate_field_names:
            if name not in all_fields:
                all_fields.append(name)

        formula_fallbacks = _formula_fallbacks(form_code, grouped)

        for field_name in all_fields:
            label_zh = get_field_label_zh(form_code, field_name)
            field_candidates = grouped.get((form_code, field_name), [])
            manual_value = manual_values.get(form_code, {}).get(field_name)
            candidate_values = {
                _text(
                    getattr(candidate, "confirmed_value", None)
                    if getattr(candidate, "confirmed_value", None) is not None
                    else getattr(candidate, "extracted_value", None)
                )
                for candidate in field_candidates
            }
            candidate_value_count = len(candidate_values)
            candidate_group = f"{form_code}.{field_name}"

            if field_candidates:
                for candidate in field_candidates:
                    confirmed_val = getattr(candidate, "confirmed_value", None)
                    extracted_val = getattr(candidate, "extracted_value", None)
                    if manual_value is not None:
                        confirmed_val = manual_value
                    page_num = getattr(candidate, "page_number", None)
                    provider = getattr(candidate, "analysis_provider", None)
                    confidence = getattr(candidate, "confidence_score", None)
                    created_at = getattr(candidate, "created_at", None)
                    direct_note = ""
                    if (
                        field_name in CALCULATED_FIELDS.get(form_code, ())
                        and _resolved_candidate_value([candidate]) is not None
                    ):
                        direct_note = (
                            f"直接確認答案（{_text(provider) or '來源未標示'}）；"
                            "未套用公式。"
                        )
                    if manual_value is not None:
                        direct_note = "；".join(
                            part for part in (direct_note, "已以手動載入值覆寫此欄位") if part
                        )
                    if candidate_value_count > 1:
                        direct_note = "；".join(
                            part
                            for part in (
                                direct_note,
                                f"同一欄位有 {candidate_value_count} 個不同候選答案；請只確認一筆正確答案。",
                            )
                            if part
                        )

                    row = [
                        form_code,
                        field_name,
                        candidate_group,
                        candidate_value_count,
                        label_zh,
                        _text(getattr(candidate, "field_status", "")),
                        _text(confirmed_val) if confirmed_val is not None else "",
                        _text(extracted_val) if extracted_val is not None else "",
                        _text(getattr(candidate, "source_text", "")),
                        _text(page_num) if page_num is not None else "",
                        _text(provider) if provider is not None else "",
                        _text(confidence) if confidence is not None else "",
                        created_at.isoformat() if isinstance(created_at, datetime) else _text(created_at),
                        direct_note,
                    ]
                    ws.append(row)
                    current_row = ws.max_row
                    for col_idx in range(1, len(headers) + 1):
                        cell = ws.cell(row=current_row, column=col_idx)
                        cell.alignment = align_left
                        cell.border = thin_border
                fallback = formula_fallbacks.get(field_name)
                if fallback is not None and _resolved_candidate_value(field_candidates) is None:
                    ws.append([
                        form_code,
                        field_name,
                        candidate_group,
                        candidate_value_count,
                        label_zh,
                        fallback.status,
                        fallback.value or "",
                        "",
                        fallback.formula,
                        "",
                        "FORMULA",
                        "",
                        generated_at.isoformat(),
                        fallback.note,
                    ])
                    current_row = ws.max_row
                    for col_idx in range(1, len(headers) + 1):
                        cell = ws.cell(row=current_row, column=col_idx)
                        cell.alignment = align_left
                        cell.border = thin_border
            elif manual_value is not None:
                row = [
                    form_code,
                    field_name,
                    candidate_group,
                    1,
                    label_zh,
                    "MANUAL_CONFIRMED",
                    _text(manual_value),
                    _text(manual_value),
                    "手動輸入後載入",
                    "",
                    "MANUAL",
                    "",
                    generated_at.isoformat(),
                    "已載入手動值；未填寫的欄位保持空白。",
                ]
                ws.append(row)
                current_row = ws.max_row
                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.alignment = align_left
                    cell.border = thin_border
            else:
                fallback = formula_fallbacks.get(field_name)
                if fallback is not None:
                    row = [
                        form_code,
                        field_name,
                        candidate_group,
                        candidate_value_count,
                        label_zh,
                        fallback.status,
                        fallback.value or "",
                        "",
                        fallback.formula,
                        "",
                        "FORMULA",
                        "",
                        generated_at.isoformat(),
                        fallback.note,
                    ]
                else:
                    row = [
                        form_code,
                        field_name,
                        candidate_group,
                        candidate_value_count,
                        label_zh,
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                    ]
                ws.append(row)
                current_row = ws.max_row
                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.alignment = align_left
                    cell.border = thin_border

    column_widths = [12, 25, 28, 16, 25, 22, 30, 30, 50, 10, 15, 12, 22, 46]
    for i, width in enumerate(column_widths, start=1):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = f"A{header_row_idx + 1}"

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
