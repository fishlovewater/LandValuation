from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

from fastapi import UploadFile
from fastapi.concurrency import run_in_threadpool
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.models import User
from app.core.exceptions import AppError, ResourceNotFoundError
from app.core.config import get_settings
from app.knowledge.models import KnowledgeDocumentRecord
from app.knowledge.source_policy import (
    EXAMPLE_REFERENCE_FILENAMES,
    is_example_reference,
)
from app.storage.paths import build_rule_source_object_key, safe_storage_filename
from app.storage.service import StorageService
from app.valuation.models import (
    FactorDefinitionRecord,
    FactorLevelRecord,
    RuleVersionRecord,
    RuleVersionSourceRecord,
)
from app.valuation.rule_packs.repository import RulePackRepository
from app.valuation.rule_packs.coverage import (
    NEW_TAIPEI_CITYWIDE_SCOPE,
    rules_overlap,
)
from app.valuation.rule_packs.schemas import (
    ExistingKnowledgeRulePackCreateRequest,
    ExistingRulePackSourceLinkRequest,
    KnowledgeSourceOptionResponse,
    RulePackImportRequest,
    RulePackImportResponse,
    RulePackAIConfirmRequest,
    RulePackAIExtractionRequest,
    RulePackAIExtractionResponse,
    RulePackAuditFinding,
    RulePackAuditResponse,
    RulePackEffectiveDateUpdate,
    RulePackManifest,
    RulePackResponse,
    RulePackSourceManifest,
    RulePackSourceResponse,
)


# A formal rule pack must show how the executable formula, factor scale and
# output layout are tied back to the supplied legal/manual/template corpus.
# PRIMARY is checked separately because it must be unique.
FORMAL_REQUIRED_SOURCE_ROLES = frozenset(
    {
        "LEGAL_BASIS",
        "NATIONAL_MANUAL",
        "LOCAL_MANUAL",
        "FACTOR_STANDARD",
        "FORM_TEMPLATE",
    }
)
from app.valuation.extraction.provider import build_document_extraction_provider
from app.valuation.rule_packs.ai_extraction import (
    GEMINI_MODEL_ID,
    build_rule_extractor,
)


ALLOWED_SOURCE_MIME_TYPES = {
    "application/pdf",
    "application/json",
    "text/csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class RulePackService:
    def __init__(
        self,
        session: AsyncSession,
        storage: StorageService,
        repository: RulePackRepository | None = None,
    ) -> None:
        self.storage = storage
        self.repository = repository or RulePackRepository(session)

    async def ai_extract_rules(
        self,
        rule_version_id: UUID,
        payload: RulePackAIExtractionRequest,
    ) -> RulePackAIExtractionResponse:
        if not payload.confirm_ai_analysis:
            raise AppError(
                "RULE_AI_CONFIRMATION_REQUIRED",
                "送交 AI 分析前必須明確確認",
                422,
            )
        rule = await self._editable_rule(rule_version_id)
        sources = await self.repository.list_sources(rule_version_id)
        selected = next(
            (
                document
                for _link, document in sources
                if payload.source_document_id is None and _link.is_primary
                or document.document_id == payload.source_document_id
            ),
            None,
        )
        if selected is None:
            raise ResourceNotFoundError("此規則版本的來源文件")

        content = await self._read_source_bytes(selected.object_key)
        source_text, extraction_provider = await self._source_text(
            content, selected.mime_type
        )
        settings = get_settings()
        source_text = source_text[: settings.ai_field_analysis_max_chars]
        result = await build_rule_extractor(settings).extract(
            source_text, list(rule.land_use_types or [])
        )
        response = RulePackAIExtractionResponse(
            rule_version_id=rule_version_id,
            source_document_id=selected.document_id,
            status="NEEDS_CONFIRMATION",
            provider=settings.ai_provider.upper(),
            model_id=GEMINI_MODEL_ID,
            prompt_version=settings.ai_rule_extraction_prompt_version,
            extraction_provider=extraction_provider,
            factor_definitions=result.factor_definitions,
            factor_levels=result.factor_levels,
            evidence=result.evidence,
            warnings=result.warnings,
        )
        summary = dict(rule.import_summary or {})
        summary["ai_extraction"] = response.model_dump(mode="json")
        rule.import_summary = summary
        rule.import_status = "AI_REVIEW_REQUIRED"
        await self.repository.save(rule)
        return response

    async def get_ai_extraction(
        self, rule_version_id: UUID
    ) -> RulePackAIExtractionResponse:
        rule = await self.repository.get(rule_version_id)
        if rule is None:
            raise ResourceNotFoundError("規則版本")
        value = (rule.import_summary or {}).get("ai_extraction")
        if not value:
            raise ResourceNotFoundError("AI 規則候選")
        return RulePackAIExtractionResponse.model_validate(value)

    async def confirm_ai_extraction(
        self,
        rule_version_id: UUID,
        payload: RulePackAIConfirmRequest,
    ) -> RulePackImportResponse:
        if not payload.confirm_import:
            raise AppError(
                "RULE_AI_IMPORT_CONFIRMATION_REQUIRED",
                "必須確認已人工檢查 AI 候選後才能匯入",
                422,
            )
        rule = await self._editable_rule(rule_version_id)
        preview = (rule.import_summary or {}).get("ai_extraction")
        if not preview:
            raise AppError("RULE_AI_PREVIEW_REQUIRED", "請先執行 AI 規則表轉換", 409)
        import_response = await self.import_rules(
            rule_version_id,
            RulePackImportRequest(
                factor_definitions=payload.factor_definitions,
                factor_levels=payload.factor_levels,
                confirm_import=True,
            ),
        )
        rule = await self._editable_rule(rule_version_id)
        summary = dict(rule.import_summary or {})
        confirmed_preview = dict(preview)
        confirmed_preview.update(
            {
                "status": "CONFIRMED",
                "factor_definitions": [
                    item.model_dump(mode="json") for item in payload.factor_definitions
                ],
                "factor_levels": [
                    item.model_dump(mode="json") for item in payload.factor_levels
                ],
                "evidence": [item.model_dump(mode="json") for item in payload.evidence],
                "confirmed_at": datetime.now(timezone.utc).isoformat(),
            }
        )
        summary["ai_extraction"] = confirmed_preview
        rule.import_summary = summary
        await self.repository.save(rule)
        return import_response

    async def _read_source_bytes(self, object_key: str) -> bytes:
        response = await self.storage.download(object_key)
        try:
            return await run_in_threadpool(response.read)
        finally:
            response.close()
            response.release_conn()

    async def _source_text(self, content: bytes, mime_type: str) -> tuple[str, str]:
        normalized = mime_type.lower()
        if normalized == "application/pdf":
            result = await build_document_extraction_provider(get_settings()).extract(content)
            return result.text, result.provider
        if normalized in {"application/json", "text/csv"}:
            return content.decode("utf-8-sig", errors="replace"), "TEXT"
        if normalized == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            def read_workbook() -> str:
                from openpyxl import load_workbook

                workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
                lines: list[str] = []
                for sheet in workbook.worksheets:
                    lines.append(f"[工作表: {sheet.title}]")
                    for row in sheet.iter_rows(values_only=True):
                        lines.append("\t".join("" if value is None else str(value) for value in row))
                workbook.close()
                return "\n".join(lines)

            return await run_in_threadpool(read_workbook), "XLSX"
        raise AppError("RULE_SOURCE_MIME_UNSUPPORTED", "此格式無法進行 AI 規則轉換", 422)

    async def upload_source(
        self,
        manifest: RulePackManifest,
        file: UploadFile,
        user: User,
    ) -> RulePackResponse:
        content_type = (file.content_type or "").lower()
        if content_type not in ALLOWED_SOURCE_MIME_TYPES:
            raise AppError(
                "RULE_SOURCE_MIME_UNSUPPORTED",
                "規則原檔只接受 PDF、JSON、CSV 或 XLSX",
                422,
            )
        version_no = await self.repository.next_version(manifest.rule_set_code)
        rule_version_id = uuid4()
        document_id = uuid4()
        filename = safe_storage_filename(file.filename)
        self._reject_example_filename(Path(file.filename or filename).name)
        object_key = build_rule_source_object_key(
            rule_version_id, document_id, version_no, filename
        )

        await file.seek(0)
        file.file.seek(0, 2)
        length = file.file.tell()
        file.file.seek(0)
        if length <= 0:
            raise AppError("EMPTY_FILE", "不可上傳空白規則檔", 422)

        uploaded = await self.storage.upload(
            object_key, file.file, length, content_type=content_type
        )
        document = KnowledgeDocumentRecord(
            document_id=document_id,
            document_code=f"RULESRC_{manifest.rule_set_code}",
            title=manifest.version_name,
            document_type=manifest.source_document_type,
            original_filename=Path(file.filename or filename).name[:255],
            mime_type=content_type,
            bucket_name=str(uploaded["bucket_name"]),
            object_key=str(uploaded["object_key"]),
            checksum_sha256=str(uploaded["checksum_sha256"]),
            file_size_bytes=int(uploaded["file_size_bytes"]),
            storage_etag=str(uploaded["etag"]),
            version_no=version_no,
            effective_from=manifest.effective_from,
            effective_to=manifest.effective_to,
            extraction_status="PENDING",
            metadata_={
                "purpose": "VALUATION_RULE_SOURCE",
                "jurisdiction_code": manifest.jurisdiction_code,
                "district_scope": dict(NEW_TAIPEI_CITYWIDE_SCOPE),
                "land_use_types": [item.value for item in manifest.land_use_types],
                "formula_code": manifest.formula_code,
                "rounding_code": manifest.rounding_code,
                "effective_date_status": manifest.effective_date_status,
            },
            created_by_user_id=user.user_id,
            publication_status="DRAFT",
        )
        rule = RuleVersionRecord(
            rule_version_id=rule_version_id,
            rule_set_code=manifest.rule_set_code,
            version_no=version_no,
            version_name=manifest.version_name,
            effective_from=manifest.effective_from,
            effective_to=manifest.effective_to,
            effective_date_status=manifest.effective_date_status,
            status="DRAFT",
            source_reference=manifest.source_reference,
            source_checksum_sha256=str(uploaded["checksum_sha256"]),
            notes=manifest.notes,
            source_document_id=document_id,
            jurisdiction_code=manifest.jurisdiction_code,
            district_scope=dict(NEW_TAIPEI_CITYWIDE_SCOPE),
            land_use_types=[item.value for item in manifest.land_use_types],
            formula_code=manifest.formula_code,
            rounding_code=manifest.rounding_code,
            import_status="SOURCE_UPLOADED",
            import_summary={},
        )
        source_link = RuleVersionSourceRecord(
            rule_version_id=rule_version_id,
            source_document_id=document_id,
            source_role=manifest.primary_source_role.value,
            source_order=1,
            is_primary=True,
            is_required=True,
            source_reference=manifest.source_reference,
            page_reference=manifest.primary_page_reference,
            notes=manifest.notes,
            created_by_user_id=user.user_id,
        )
        document.metadata_.update(
            {
                "source_role": manifest.primary_source_role.value,
                "source_reference": manifest.source_reference,
                "page_reference": manifest.primary_page_reference,
                "is_primary": True,
                "is_required": True,
            }
        )
        try:
            created = await self.repository.create_source(document, rule, source_link)
        except Exception:
            await self.storage.delete(object_key)
            raise
        return await self._response(created, document)

    async def create_from_existing_source(
        self,
        payload: ExistingKnowledgeRulePackCreateRequest,
        user: User,
    ) -> RulePackResponse:
        document = await self.repository.get_source(payload.source_document_id)
        if document is None:
            raise ResourceNotFoundError("知識來源文件")
        self._reject_example_source(document)
        if document.mime_type.lower() not in ALLOWED_SOURCE_MIME_TYPES:
            raise AppError(
                "RULE_SOURCE_MIME_UNSUPPORTED",
                "既有規則來源只接受 PDF、JSON、CSV 或 XLSX",
                422,
            )
        if not document.object_key.startswith("knowledge/"):
            raise AppError(
                "RULE_SOURCE_KEY_INVALID",
                "規則來源必須是 knowledge/ 下的 MinIO 物件",
                422,
            )
        if not await self.storage.object_exists(document.object_key):
            raise AppError(
                "RULE_SOURCE_OBJECT_MISSING",
                "知識文件 metadata 存在，但 MinIO 原始物件不存在",
                409,
            )

        version_no = await self.repository.next_version(payload.rule_set_code)
        rule_version_id = uuid4()
        rule = RuleVersionRecord(
            rule_version_id=rule_version_id,
            rule_set_code=payload.rule_set_code,
            version_no=version_no,
            version_name=payload.version_name,
            effective_from=payload.effective_from,
            effective_to=payload.effective_to,
            effective_date_status=payload.effective_date_status,
            status="DRAFT",
            source_reference=payload.source_reference,
            source_checksum_sha256=document.checksum_sha256,
            notes=payload.notes,
            source_document_id=document.document_id,
            jurisdiction_code=payload.jurisdiction_code,
            district_scope=dict(NEW_TAIPEI_CITYWIDE_SCOPE),
            land_use_types=[item.value for item in payload.land_use_types],
            formula_code=payload.formula_code,
            rounding_code=payload.rounding_code,
            import_status="SOURCE_UPLOADED",
            import_summary={},
        )
        source_link = RuleVersionSourceRecord(
            rule_version_id=rule_version_id,
            source_document_id=document.document_id,
            source_role="PRIMARY",
            source_order=1,
            is_primary=True,
            is_required=True,
            source_reference=payload.source_reference,
            page_reference=payload.primary_page_reference,
            notes=payload.notes,
            created_by_user_id=user.user_id,
        )
        created = await self.repository.create_rule_with_source_link(rule, source_link)
        return await self._response(created, document)

    async def add_source(
        self,
        rule_version_id: UUID,
        manifest: RulePackSourceManifest,
        file: UploadFile,
        user: User,
    ) -> RulePackSourceResponse:
        rule = await self._editable_rule(rule_version_id)
        content_type = (file.content_type or "").lower()
        if content_type not in ALLOWED_SOURCE_MIME_TYPES:
            raise AppError(
                "RULE_SOURCE_MIME_UNSUPPORTED",
                "規則原檔只接受 PDF、JSON、CSV 或 XLSX",
                422,
            )

        source_order = await self.repository.next_source_order(rule_version_id)
        document_id = uuid4()
        filename = safe_storage_filename(file.filename)
        self._reject_example_filename(Path(file.filename or filename).name)
        object_key = build_rule_source_object_key(
            rule_version_id, document_id, rule.version_no, filename
        )

        await file.seek(0)
        file.file.seek(0, 2)
        length = file.file.tell()
        file.file.seek(0)
        if length <= 0:
            raise AppError("EMPTY_FILE", "不可上傳空白規則檔", 422)

        uploaded = await self.storage.upload(
            object_key, file.file, length, content_type=content_type
        )
        existing_sources = await self.repository.list_sources(rule_version_id)
        duplicate = next(
            (
                document
                for _, document in existing_sources
                if document.checksum_sha256 == str(uploaded["checksum_sha256"])
            ),
            None,
        )
        if duplicate is not None:
            await self.storage.delete(object_key)
            raise AppError(
                "RULE_SOURCE_DUPLICATE",
                f"相同內容已存在於來源文件：{duplicate.original_filename}",
                409,
            )

        document = KnowledgeDocumentRecord(
            document_id=document_id,
            document_code=f"RULESRC_{rule.rule_set_code}_{source_order}",
            title=manifest.title,
            document_type=manifest.source_document_type,
            original_filename=Path(file.filename or filename).name[:255],
            mime_type=content_type,
            bucket_name=str(uploaded["bucket_name"]),
            object_key=str(uploaded["object_key"]),
            checksum_sha256=str(uploaded["checksum_sha256"]),
            file_size_bytes=int(uploaded["file_size_bytes"]),
            storage_etag=str(uploaded["etag"]),
            version_no=rule.version_no,
            effective_from=None,
            effective_to=None,
            extraction_status="PENDING",
            metadata_={
                "purpose": "VALUATION_RULE_SOURCE",
                "rule_version_id": str(rule_version_id),
                "source_role": manifest.source_role.value,
                "source_reference": manifest.source_reference,
                "page_reference": manifest.page_reference,
                "is_primary": False,
                "is_required": manifest.is_required,
            },
            created_by_user_id=user.user_id,
            publication_status="DRAFT",
        )
        source_link = RuleVersionSourceRecord(
            rule_version_id=rule_version_id,
            source_document_id=document_id,
            source_role=manifest.source_role.value,
            source_order=source_order,
            is_primary=False,
            is_required=manifest.is_required,
            source_reference=manifest.source_reference,
            page_reference=manifest.page_reference,
            notes=manifest.notes,
            created_by_user_id=user.user_id,
        )
        try:
            created = await self.repository.add_source(document, source_link)
        except Exception:
            await self.storage.delete(object_key)
            raise
        return self._source_response(created, document)

    async def link_existing_source(
        self,
        rule_version_id: UUID,
        payload: ExistingRulePackSourceLinkRequest,
        user: User,
    ) -> RulePackSourceResponse:
        await self._editable_rule(rule_version_id)
        document = await self.repository.get_source(payload.source_document_id)
        if document is None:
            raise ResourceNotFoundError("知識來源文件")
        self._reject_example_source(document)
        if not document.object_key.startswith("knowledge/"):
            raise AppError(
                "RULE_SOURCE_OBJECT_KEY_INVALID",
                "規則來源必須是 knowledge/ 路徑中的知識文件",
                409,
            )
        if not await self.storage.object_exists(document.object_key):
            raise AppError(
                "RULE_SOURCE_OBJECT_MISSING",
                f"MinIO 找不到知識來源：{document.original_filename}",
                409,
            )

        existing_sources = await self.repository.list_sources(rule_version_id)
        if any(
            source_document.document_id == payload.source_document_id
            for _, source_document in existing_sources
        ):
            raise AppError("RULE_SOURCE_ALREADY_LINKED", "此知識文件已關聯規則版本", 409)

        source_order = await self.repository.next_source_order(rule_version_id)
        source_link = RuleVersionSourceRecord(
            rule_version_id=rule_version_id,
            source_document_id=document.document_id,
            source_role=payload.source_role.value,
            source_order=source_order,
            is_primary=False,
            is_required=payload.is_required,
            source_reference=payload.source_reference,
            page_reference=payload.page_reference,
            notes=payload.notes,
            created_by_user_id=user.user_id,
        )
        created = await self.repository.add_source_link(source_link)
        return self._source_response(created, document)

    async def confirm_effective_date(
        self,
        rule_version_id: UUID,
        payload: RulePackEffectiveDateUpdate,
    ) -> RulePackResponse:
        rule = await self._editable_rule(rule_version_id)
        rule.effective_from = payload.effective_from
        rule.effective_to = payload.effective_to
        rule.effective_date_status = "CONFIRMED"
        saved = await self.repository.save(rule)
        source = (
            await self.repository.get_source(rule.source_document_id)
            if rule.source_document_id
            else None
        )
        return await self._response(saved, source)

    async def import_rules(
        self,
        rule_version_id: UUID,
        payload: RulePackImportRequest,
    ) -> RulePackImportResponse:
        rule = await self._editable_rule(rule_version_id)
        definitions = payload.factor_definitions
        levels = payload.factor_levels
        codes = [item.factor_code for item in definitions]
        orders = [item.display_order for item in definitions]
        if len(codes) != len(set(codes)) or len(orders) != len(set(orders)):
            raise AppError(
                "RULE_FACTOR_DUPLICATE",
                "因素代碼及顯示順序不可重複",
                422,
            )
        definition_by_code = {item.factor_code: item for item in definitions}
        unknown = sorted({item.factor_code for item in levels} - set(codes))
        if unknown:
            raise AppError(
                "RULE_FACTOR_NOT_DEFINED",
                f"級距引用未定義因素：{', '.join(unknown)}",
                422,
            )
        level_keys = [
            (item.factor_code, item.land_use_type.value, item.level_code)
            for item in levels
        ]
        if len(level_keys) != len(set(level_keys)):
            raise AppError("RULE_LEVEL_DUPLICATE", "同一因素、用途及級距代碼不可重複", 422)

        declared = set(rule.land_use_types or [])
        actual = {item.land_use_type.value for item in levels}
        undeclared = sorted(actual - declared)
        missing = sorted(declared - actual)
        if undeclared or missing:
            message = []
            if undeclared:
                message.append(f"未宣告用途：{', '.join(undeclared)}")
            if missing:
                message.append(f"缺少用途級距：{', '.join(missing)}")
            raise AppError("RULE_LAND_USE_COVERAGE_INVALID", "；".join(message), 422)

        grouped_levels = {}
        for item in levels:
            grouped_levels.setdefault(
                (item.factor_code, item.land_use_type.value), []
            ).append(item)
        missing_factor_coverage = [
            f"{factor_code}:{land_use}"
            for factor_code in codes
            for land_use in sorted(declared)
            if (factor_code, land_use) not in grouped_levels
        ]
        if missing_factor_coverage:
            raise AppError(
                "RULE_FACTOR_LAND_USE_COVERAGE_INCOMPLETE",
                "每個正式因素都必須對已宣告的每一種土地用途提供級距",
                422,
                {"missing": missing_factor_coverage},
            )
        for (factor_code, land_use), items in grouped_levels.items():
            if len(items) < 2:
                raise AppError(
                    "RULE_LEVELS_INCOMPLETE",
                    f"因素 {factor_code}／{land_use} 至少需要兩個級距",
                    422,
                )
            orders_for_factor = [item.sort_order for item in items]
            names_for_factor = [item.level_name for item in items]
            if len(orders_for_factor) != len(set(orders_for_factor)):
                raise AppError(
                    "RULE_LEVEL_ORDER_DUPLICATE",
                    f"因素 {factor_code}／{land_use} 的級距順序不可重複",
                    422,
                )
            if len(names_for_factor) != len(set(names_for_factor)):
                raise AppError(
                    "RULE_LEVEL_NAME_DUPLICATE",
                    f"因素 {factor_code}／{land_use} 的級距名稱不可重複",
                    422,
                )
            bounds = {item.maximum_impact_rate for item in items}
            if len(bounds) != 1:
                raise AppError(
                    "RULE_FACTOR_BOUND_INCONSISTENT",
                    f"因素 {factor_code}／{land_use} 的最大影響範圍必須一致",
                    422,
                )
            rates = [item.suggested_rate for item in items]
            if max(rates) - min(rates) > next(iter(bounds)):
                raise AppError(
                    "RULE_FACTOR_SPAN_OUT_OF_RANGE",
                    f"因素 {factor_code}／{land_use} 的級距修正跨度超出最大影響範圍",
                    422,
                )

        existing = await self.repository.matching_definitions(set(codes), set(orders))
        existing_by_code = {item.factor_code: item for item in existing}
        for item in existing:
            incoming = definition_by_code.get(item.factor_code)
            if incoming is None:
                raise AppError(
                    "RULE_FACTOR_ORDER_CONFLICT",
                    f"顯示順序 {item.display_order} 已由因素 {item.factor_code} 使用",
                    409,
                )
            comparable = (
                item.factor_name,
                item.factor_category,
                item.data_type,
                item.unit,
                item.display_order,
            )
            expected = (
                incoming.factor_name,
                incoming.factor_category,
                incoming.data_type,
                incoming.unit,
                incoming.display_order,
            )
            if comparable != expected:
                raise AppError(
                    "RULE_FACTOR_DEFINITION_CONFLICT",
                    f"既有因素 {item.factor_code} 定義與匯入內容不同",
                    409,
                )

        if not payload.confirm_import:
            return RulePackImportResponse(
                rule_version_id=rule_version_id,
                confirmed=False,
                factor_definition_count=len(definitions),
                factor_level_count=len(levels),
                land_use_types=sorted(actual),
                status="PREVIEW_VALID",
                warnings=["尚未寫入；confirm_import=true 後才會保存"],
            )

        new_definitions = [
            FactorDefinitionRecord(
                factor_code=item.factor_code,
                factor_name=item.factor_name,
                factor_category=item.factor_category,
                data_type=item.data_type,
                unit=item.unit,
                display_order=item.display_order,
                is_active=True,
            )
            for item in definitions
            if item.factor_code not in existing_by_code
        ]
        all_records = list(existing_by_code.values()) + new_definitions
        record_by_code = {item.factor_code: item for item in all_records}
        # New definition UUIDs are assigned before flush by the model default only on insert;
        # assign explicitly so level objects can safely reference them in the same transaction.
        for item in new_definitions:
            if item.factor_definition_id is None:
                item.factor_definition_id = uuid4()
        level_records = [
            FactorLevelRecord(
                factor_definition_id=record_by_code[item.factor_code].factor_definition_id,
                rule_version_id=rule_version_id,
                land_use_type=item.land_use_type.value,
                level_code=item.level_code,
                level_name=item.level_name,
                range_min=item.range_min,
                range_max=item.range_max,
                qualitative_value=item.qualitative_value,
                suggested_rate=item.suggested_rate,
                maximum_impact_rate=item.maximum_impact_rate,
                sort_order=item.sort_order,
            )
            for item in levels
        ]
        await self.repository.replace_levels(rule_version_id, new_definitions, level_records)
        rule.import_status = "IMPORTED"
        ai_extraction = (rule.import_summary or {}).get("ai_extraction")
        rule.import_summary = {
            "factor_definition_count": len(definitions),
            "factor_level_count": len(levels),
            "land_use_types": sorted(actual),
        }
        if ai_extraction:
            rule.import_summary["ai_extraction"] = ai_extraction
        await self.repository.save(rule)
        return RulePackImportResponse(
            rule_version_id=rule_version_id,
            confirmed=True,
            factor_definition_count=len(definitions),
            factor_level_count=len(levels),
            land_use_types=sorted(actual),
            status="IMPORTED",
        )

    async def publish(
        self, rule_version_id: UUID, confirm_publish: bool, user: User
    ) -> RulePackResponse:
        if not confirm_publish:
            raise AppError("RULE_PUBLISH_CONFIRMATION_REQUIRED", "發布前必須明確確認", 422)
        rule = await self._editable_rule(rule_version_id)
        if rule.import_status != "IMPORTED":
            raise AppError("RULE_IMPORT_REQUIRED", "必須先完成規則級距匯入", 409)
        if rule.effective_date_status != "CONFIRMED" or rule.effective_from is None:
            raise AppError(
                "RULE_EFFECTIVE_DATE_REQUIRED",
                "發布前必須由使用者確認規則適用起日",
                409,
            )
        sources = await self.repository.list_sources(rule_version_id)
        if not sources:
            raise AppError("RULE_SOURCE_REQUIRED", "規則版本缺少原始來源文件", 409)
        primary_sources = [item for item in sources if item[0].is_primary]
        if (
            len(primary_sources) != 1
            or rule.source_document_id != primary_sources[0][1].document_id
        ):
            raise AppError("RULE_PRIMARY_SOURCE_INVALID", "規則版本主要來源關聯不一致", 409)
        source_roles = {item[0].source_role for item in sources}
        missing_roles = sorted(FORMAL_REQUIRED_SOURCE_ROLES - source_roles)
        if missing_roles:
            raise AppError(
                "RULE_FORMAL_SOURCES_INCOMPLETE",
                "正式規則尚未完整關聯法規、作業手冊、製作手冊、因素準則及表單範本",
                409,
                {"missing_source_roles": missing_roles},
            )
        for _, source_document in sources:
            self._reject_example_source(source_document)
            if not await self.storage.object_exists(source_document.object_key):
                raise AppError(
                    "RULE_SOURCE_OBJECT_MISSING",
                    f"MinIO 找不到已關聯規則來源：{source_document.original_filename}",
                    409,
                )
        level_count = await self.repository.level_count(rule_version_id)
        coverage = await self.repository.level_land_uses(rule_version_id)
        if level_count == 0 or coverage != set(rule.land_use_types or []):
            raise AppError("RULE_LEVELS_INCOMPLETE", "規則級距或土地用途覆蓋不完整", 409)

        overlapping = [
            item
            for item in await self.repository.list_published()
            if item.rule_version_id != rule.rule_version_id
            and rules_overlap(rule, item)
        ]
        if overlapping:
            raise AppError(
                "RULE_COVERAGE_OVERLAP",
                "相同行政區、土地用途及生效期間已有正式規則，禁止產生重疊版本",
                409,
                {
                    "overlapping_rule_version_ids": [
                        str(item.rule_version_id) for item in overlapping
                    ]
                },
            )

        now = datetime.now(timezone.utc)
        rule.status = "PUBLISHED"
        rule.import_status = "VERIFIED"
        rule.verified_by_user_id = user.user_id
        rule.verified_at = now
        for _, source_document in sources:
            source_document.publication_status = "PUBLISHED"
            # A source is already read/imported and manually verified before a
            # rule pack can be published.  Keep its lifecycle in sync with the
            # published rule so the same fixed rule can be used by formal
            # calculation and review handoff.  Leaving this at PENDING made
            # every otherwise-valid published rule fail submission preflight.
            source_document.extraction_status = "COMPLETED"
            source_document.approved_by_user_id = user.user_id
            source_document.approved_at = now
            source_document.updated_at = now
        saved = await self.repository.save(rule)
        return await self._response(saved, primary_sources[0][1])

    async def get(self, rule_version_id: UUID) -> RulePackResponse:
        rule = await self.repository.get(rule_version_id)
        if rule is None or rule.jurisdiction_code != "NEW_TAIPEI_CITY":
            raise ResourceNotFoundError("規則版本")
        source = (
            await self.repository.get_source(rule.source_document_id)
            if rule.source_document_id
            else None
        )
        return await self._response(rule, source)

    async def list(self) -> list[RulePackResponse]:
        responses = []
        for rule in await self.repository.list_all():
            source = (
                await self.repository.get_source(rule.source_document_id)
                if rule.source_document_id
                else None
            )
            responses.append(await self._response(rule, source))
        return responses

    async def formal_audit(self, rule_version_id: UUID) -> RulePackAuditResponse:
        rule = await self.repository.get(rule_version_id)
        if rule is None or rule.jurisdiction_code != "NEW_TAIPEI_CITY":
            raise ResourceNotFoundError("規則版本")
        sources = await self.repository.list_sources(rule_version_id)
        level_rows = await self.repository.audit_levels(rule_version_id)
        findings: list[RulePackAuditFinding] = []

        def error(code: str, message: str, factor=None, land_use=None) -> None:
            findings.append(
                RulePackAuditFinding(
                    code=code,
                    severity="ERROR",
                    message=message,
                    factor_code=factor,
                    land_use_type=land_use,
                )
            )

        def warning(code: str, message: str) -> None:
            findings.append(
                RulePackAuditFinding(
                    code=code,
                    severity="WARNING",
                    message=message,
                )
            )

        if rule.status != "PUBLISHED" or rule.import_status != "VERIFIED":
            error("RULE_NOT_PUBLISHED_AND_VERIFIED", "規則尚未完成匯入及人工發布")
        if rule.effective_date_status != "CONFIRMED" or rule.effective_from is None:
            error("RULE_EFFECTIVE_DATE_UNCONFIRMED", "規則適用起日尚未確認")
        if rule.formula_code != "NTPC_COMPARISON_V1":
            error("RULE_FORMULA_UNSUPPORTED", "正式比較法公式代碼不支援")
        if rule.rounding_code != "NTPC_LAND_PRICE_V1":
            error("RULE_ROUNDING_UNSUPPORTED", "正式地價尾數規則代碼不支援")
        if rule.district_scope != NEW_TAIPEI_CITYWIDE_SCOPE:
            error("RULE_SCOPE_NOT_CITYWIDE", "本系統正式共同規則必須適用新北市29區全市")
        if not sources:
            error("RULE_SOURCES_MISSING", "規則未關聯任何正式來源文件")
        primary = [item for item in sources if item[0].is_primary]
        if len(primary) != 1:
            error("RULE_PRIMARY_SOURCE_INVALID", "規則必須且只能有一份主要來源")
        source_roles = {item[0].source_role for item in sources}
        for role in sorted(FORMAL_REQUIRED_SOURCE_ROLES - source_roles):
            error(
                "RULE_FORMAL_SOURCE_ROLE_MISSING",
                f"正式規則缺少必要來源角色：{role}",
            )
        for link, document in sources:
            if is_example_reference(document):
                error(
                    "EXAMPLE_SOURCE_NOT_FORMAL",
                    f"範例文件不可作正式規則來源：{document.original_filename}",
                )
            if not await self.storage.object_exists(document.object_key):
                error(
                    "RULE_SOURCE_OBJECT_MISSING",
                    f"MinIO 找不到規則來源：{document.original_filename}",
                )
            if not link.page_reference:
                warning(
                    "RULE_SOURCE_PAGE_REFERENCE_MISSING",
                    f"來源尚未標示採用頁次：{document.original_filename}",
                )

        grouped = {}
        definitions = {}
        for definition, level in level_rows:
            definitions[definition.factor_code] = definition
            grouped.setdefault(
                (definition.factor_code, level.land_use_type), []
            ).append(level)
        declared = set(rule.land_use_types or [])
        for factor_code in definitions:
            for land_use in declared:
                items = grouped.get((factor_code, land_use), [])
                if len(items) < 2:
                    error(
                        "RULE_LEVELS_INCOMPLETE",
                        "每個正式因素與土地用途至少需要兩個級距",
                        factor_code,
                        land_use,
                    )
                    continue
                bounds = {item.maximum_impact_rate for item in items}
                if len(bounds) != 1:
                    error(
                        "RULE_FACTOR_BOUND_INCONSISTENT",
                        "同一因素的最大影響範圍不一致",
                        factor_code,
                        land_use,
                    )
                    continue
                rates = [item.suggested_rate for item in items]
                if max(rates) - min(rates) > next(iter(bounds)):
                    error(
                        "RULE_FACTOR_SPAN_OUT_OF_RANGE",
                        "級距修正跨度超出最大影響範圍",
                        factor_code,
                        land_use,
                    )
        if not definitions:
            error("RULE_FACTOR_DEFINITIONS_MISSING", "規則未提供任何正式因素級距")

        from app.valuation.report_packages.factor_catalog import (
            INDIVIDUAL_FACTOR_CODES,
            TEMPLATE_FACTOR_CODES,
        )

        if "COMMERCIAL" in declared:
            required = TEMPLATE_FACTOR_CODES | INDIVIDUAL_FACTOR_CODES
            missing = sorted(required - set(definitions))
            for factor_code in missing:
                error(
                    "COMMERCIAL_REPORT_FACTOR_MISSING",
                    "商業用地六頁查估書缺少必要因素",
                    factor_code,
                    "COMMERCIAL",
                )
        all_land_uses = {
            "RESIDENTIAL",
            "COMMERCIAL",
            "INDUSTRIAL",
            "AGRICULTURAL",
            "OTHER",
        }
        if declared != all_land_uses:
            warning(
                "RULE_PACK_NOT_ALL_LAND_USES",
                "此單一規則版本未涵蓋全部五種土地用途；全市覆蓋仍需由 coverage matrix 確認",
            )

        error_count = sum(item.severity == "ERROR" for item in findings)
        warning_count = sum(item.severity == "WARNING" for item in findings)
        return RulePackAuditResponse(
            rule_version_id=rule.rule_version_id,
            rule_set_code=rule.rule_set_code,
            version_no=rule.version_no,
            status=rule.status,
            formal_calculation_eligible=error_count == 0,
            jurisdiction_code=rule.jurisdiction_code,
            district_scope=rule.district_scope,
            land_use_types=list(rule.land_use_types or []),
            source_count=len(sources),
            factor_definition_count=len(definitions),
            factor_level_count=len(level_rows),
            error_count=error_count,
            warning_count=warning_count,
            findings=findings,
        )

    async def source_document(
        self, rule_version_id: UUID
    ) -> KnowledgeDocumentRecord:
        rule = await self.repository.get(rule_version_id)
        if rule is None or rule.source_document_id is None:
            raise ResourceNotFoundError("規則原始檔")
        source = await self.repository.get_source(rule.source_document_id)
        if source is None:
            raise ResourceNotFoundError("規則原始檔")
        return source

    async def list_source_documents(
        self, rule_version_id: UUID
    ) -> list[RulePackSourceResponse]:
        rule = await self.repository.get(rule_version_id)
        if rule is None or rule.jurisdiction_code != "NEW_TAIPEI_CITY":
            raise ResourceNotFoundError("規則版本")
        return [
            self._source_response(source_link, document)
            for source_link, document in await self.repository.list_sources(rule_version_id)
        ]

    async def list_available_knowledge_sources(
        self,
    ) -> list[KnowledgeSourceOptionResponse]:
        responses = []
        for document in await self.repository.list_available_knowledge_sources():
            responses.append(
                KnowledgeSourceOptionResponse(
                    source_document_id=document.document_id,
                    document_type=document.document_type,
                    title=document.title,
                    original_filename=document.original_filename,
                    mime_type=document.mime_type,
                    bucket_name=document.bucket_name,
                    object_key=document.object_key,
                    checksum_sha256=document.checksum_sha256,
                    version_no=document.version_no,
                    effective_from=document.effective_from,
                    effective_to=document.effective_to,
                    publication_status=document.publication_status,
                    object_exists=await self.storage.object_exists(document.object_key),
                    formal_rule_eligible=not is_example_reference(document),
                    source_usage=(
                        "EXAMPLE_REFERENCE"
                        if is_example_reference(document)
                        else "FORMAL_SOURCE_CANDIDATE"
                    ),
                )
            )
        return responses

    async def rule_source_document(
        self, rule_version_id: UUID, rule_version_source_id: UUID
    ) -> KnowledgeDocumentRecord:
        rule = await self.repository.get(rule_version_id)
        if rule is None or rule.jurisdiction_code != "NEW_TAIPEI_CITY":
            raise ResourceNotFoundError("規則版本")
        row = await self.repository.get_rule_source(
            rule_version_id, rule_version_source_id
        )
        if row is None:
            raise ResourceNotFoundError("規則來源文件")
        return row[1]

    async def _editable_rule(self, rule_version_id: UUID) -> RuleVersionRecord:
        rule = await self.repository.get(rule_version_id)
        if rule is None or rule.jurisdiction_code != "NEW_TAIPEI_CITY":
            raise ResourceNotFoundError("規則版本")
        if rule.status != "DRAFT":
            raise AppError("RULE_VERSION_LOCKED", "已發布規則不可覆寫", 409)
        return rule

    async def _response(
        self,
        rule: RuleVersionRecord,
        source: KnowledgeDocumentRecord | None,
    ) -> RulePackResponse:
        source_rows = await self.repository.list_sources(rule.rule_version_id)
        source_responses = [
            self._source_response(source_link, document)
            for source_link, document in source_rows
        ]
        return RulePackResponse(
            rule_version_id=rule.rule_version_id,
            rule_set_code=rule.rule_set_code,
            version_no=rule.version_no,
            version_name=rule.version_name,
            effective_from=rule.effective_from,
            effective_to=rule.effective_to,
            effective_date_status=rule.effective_date_status,
            status=rule.status,
            import_status=rule.import_status,
            jurisdiction_code=rule.jurisdiction_code,
            district_scope=rule.district_scope,
            land_use_types=list(rule.land_use_types or []),
            formula_code=rule.formula_code,
            rounding_code=rule.rounding_code,
            source_document_id=rule.source_document_id,
            source_filename=source.original_filename if source else None,
            source_object_key=source.object_key if source else None,
            source_checksum_sha256=source.checksum_sha256 if source else None,
            source_count=len(source_responses),
            required_source_count=sum(item.is_required for item in source_responses),
            sources=source_responses,
            import_summary=dict(rule.import_summary or {}),
            verified_by_user_id=rule.verified_by_user_id,
            verified_at=rule.verified_at,
            created_at=rule.created_at,
        )

    @staticmethod
    def _source_response(
        source_link: RuleVersionSourceRecord,
        document: KnowledgeDocumentRecord,
    ) -> RulePackSourceResponse:
        metadata = dict(document.metadata_ or {})
        return RulePackSourceResponse(
            rule_version_source_id=source_link.rule_version_source_id,
            rule_version_id=source_link.rule_version_id,
            source_document_id=document.document_id,
            source_role=source_link.source_role,
            source_order=source_link.source_order,
            is_primary=source_link.is_primary,
            is_required=source_link.is_required,
            source_document_type=document.document_type,
            title=document.title,
            original_filename=document.original_filename,
            mime_type=document.mime_type,
            bucket_name=document.bucket_name,
            object_key=document.object_key,
            checksum_sha256=document.checksum_sha256,
            file_size_bytes=document.file_size_bytes,
            storage_etag=document.storage_etag,
            source_reference=source_link.source_reference
            or metadata.get("source_reference"),
            page_reference=source_link.page_reference,
            notes=source_link.notes,
            publication_status=document.publication_status,
            created_at=document.created_at,
        )

    @staticmethod
    def _reject_example_filename(filename: str) -> None:
        if Path(filename).name in EXAMPLE_REFERENCE_FILENAMES:
            raise AppError(
                "EXAMPLE_REFERENCE_NOT_RULE_SOURCE",
                "金山區評價基準明細表僅供學習參考，不得作為正式規則來源",
                422,
            )

    @staticmethod
    def _reject_example_source(document: KnowledgeDocumentRecord) -> None:
        if is_example_reference(document):
            raise AppError(
                "EXAMPLE_REFERENCE_NOT_RULE_SOURCE",
                "範例參考文件不得建立、關聯或發布為正式規則",
                422,
                {"source_document_id": str(document.document_id)},
            )
