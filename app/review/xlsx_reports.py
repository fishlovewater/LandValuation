"""Pure `.xlsx` renderer for an immutable Review risk report.

The workbook is rendered from a single `ReviewReport` snapshot. It never reads
current database rows, never contains macros, and never exposes bucket names,
object keys, or storage URLs. No formula computes an appraisal value.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.review.report_presentation import (
    CORRECTION_STATUS_LABELS,
    DOCUMENT_TYPE_LABELS,
    EVENT_LABELS,
    FINDING_TYPE_LABELS,
    OUTCOME_LABELS,
    REVIEW_STATUS_LABELS,
    RISK_LABELS,
    SEVERITY_LABELS,
    SOURCE_LABELS,
    TRIAGE_LABELS,
    URGENCY_LABELS,
    decision_label,
    district_label,
    label,
    local_time,
    readable_text,
    skipped_rule_text,
)
from app.review.reports import ReviewReport


SHEET_SUMMARY = "案件摘要"
SHEET_DOCUMENTS = "審查文件"
SHEET_FINDINGS = "疑點與修正要求"
SHEET_RECHECK = "新版重檢結果"
SHEET_HISTORY = "審查歷程"

FINDING_HEADERS = [
    "疑點類型",
    "嚴重度",
    "人工判定",
    "疑點說明",
    "原報告頁碼",
    "原報告內容",
    "原報告數值",
    "法規依據",
    "查核證據",
    "建議修正方向",
    "AI 輔助說明",
]

RECHECK_HEADERS = [
    "修正通知編號",
    "嚴重度",
    "問題摘要",
    "要求修正內容",
    "重檢結果",
    "重檢時間",
]

HISTORY_HEADERS = ["事件", "發生時間", "說明"]


def _finding_status_label(status: str) -> str:
    return (
        decision_label(status)
        if status not in TRIAGE_LABELS
        else TRIAGE_LABELS[status]
    )


def _autosize(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_headers(sheet, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"
    )


def build_review_xlsx(report: ReviewReport) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    _build_summary(workbook.create_sheet(SHEET_SUMMARY), report)
    _build_documents(workbook.create_sheet(SHEET_DOCUMENTS), report)
    _build_findings(workbook.create_sheet(SHEET_FINDINGS), report)
    _build_recheck(workbook.create_sheet(SHEET_RECHECK), report)
    _build_history(workbook.create_sheet(SHEET_HISTORY), report)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _build_summary(sheet, report: ReviewReport) -> None:
    urgency = report.urgency
    provenance = report.input_provenance
    provenance_source = label(SOURCE_LABELS, provenance.source) if provenance else ""
    rows = [
        ("案件編號", report.case.case_no),
        ("案件名稱", report.case.case_title),
        ("行政區", district_label(report.case.district_code)),
        ("價格基準日", report.case.valuation_base_date),
        ("審查狀態", label(REVIEW_STATUS_LABELS, report.review_status)),
        ("審查開始", local_time(report.run.started_at)),
        ("審查完成", local_time(report.run.completed_at)),
        ("審查依據來源", provenance_source),
        (
            "審查輸入版本",
            f"v{provenance.version_no}"
            if provenance and provenance.version_no is not None
            else "",
        ),
        ("內容風險等級", label(RISK_LABELS, report.risk_summary.overall_risk_level)),
        ("高風險疑點數", report.risk_summary.high_count),
        ("中風險疑點數", report.risk_summary.medium_count),
        ("低風險疑點數", report.risk_summary.low_count),
        ("缺件數", report.missing_item_count),
        ("期限緊急度", URGENCY_LABELS.get(urgency.level, urgency.level) if urgency else "未設定"),
        ("剩餘天數", urgency.remaining_days if urgency else ""),
        ("修正期限", local_time(urgency.due_at) if urgency else ""),
        ("修正通知次數", len(report.correction_requests)),
    ]
    if report.review_coverage is not None:
        rows.extend(
            [
                ("適用檢核規則總數", report.review_coverage.total_rule_count),
                ("已執行檢核規則", report.review_coverage.executed_rule_count),
                ("未執行檢核規則", report.review_coverage.skipped_rule_count),
                (
                    "檢核覆蓋說明",
                    "未執行不代表通過；資料不足的規則會保留在未執行清單。",
                ),
                (
                    "未執行規則",
                    "\n".join(
                        skipped_rule_text(item)
                        for item in report.review_coverage.skipped_rules
                    ),
                ),
            ]
        )
    sheet.append(["項目", "內容"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row_label, value in rows:
        sheet.append([row_label, value if value is not None else ""])
    for row in sheet.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.font = Font(bold=True)
    _autosize(sheet, [22, 60])


def _build_documents(sheet, report: ReviewReport) -> None:
    provenance = report.input_provenance
    sheet.append(["文件名稱", "文件類型", "文件版本"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    if provenance:
        for document in provenance.documents:
            sheet.append(
                [
                    document.original_filename or "未命名文件",
                    label(DOCUMENT_TYPE_LABELS, document.document_type),
                    f"v{document.version_no}",
                ]
            )
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    _autosize(sheet, [38, 28, 14])


def _build_findings(sheet, report: ReviewReport) -> None:
    _write_headers(sheet, FINDING_HEADERS)
    requested_by_finding = {
        item.finding_id: item
        for request in report.correction_requests
        for item in request.items
    }
    for finding in report.findings:
        requested = requested_by_finding.get(finding.finding_id)
        page = ""
        for evidence in finding.source_evidence:
            if isinstance(evidence, dict) and evidence.get("page"):
                page = evidence["page"]
                break
        sheet.append(
            [
                label(FINDING_TYPE_LABELS, finding.finding_type),
                label(SEVERITY_LABELS, finding.severity),
                _finding_status_label(finding.status),
                f"{finding.title}：{finding.description}",
                page,
                finding.reported_text or "",
                finding.reported_value or "",
                readable_text(finding.legal_basis),
                readable_text(finding.source_evidence),
                requested.requested_correction if requested else "",
                finding.ai_assessment.reasoning_summary or "",
            ]
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _autosize(sheet, [20, 10, 20, 40, 12, 24, 14, 32, 40, 32, 40])
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(FINDING_HEADERS))}{max(sheet.max_row, 1)}"
    )


def _build_recheck(sheet, report: ReviewReport) -> None:
    _write_headers(sheet, RECHECK_HEADERS)
    for request in report.correction_requests:
        for item in request.items:
            sheet.append(
                [
                    request.request_no,
                    label(SEVERITY_LABELS, item.severity),
                    item.issue_summary,
                    item.requested_correction,
                    label(OUTCOME_LABELS, item.recheck_outcome),
                    local_time(request.rechecked_at),
                ]
            )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _autosize(sheet, [16, 10, 36, 36, 14, 18])
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(RECHECK_HEADERS))}{max(sheet.max_row, 1)}"
    )


def _build_history(sheet, report: ReviewReport) -> None:
    _write_headers(sheet, HISTORY_HEADERS)
    for event in report.history:
        sheet.append(
            [label(EVENT_LABELS, event.event_type), local_time(event.occurred_at), event.reason or ""]
        )
    for request in report.correction_requests:
        sheet.append(
            [
                f"第 {request.request_no} 次修正通知（{label(CORRECTION_STATUS_LABELS, request.status)}）",
                local_time(request.sent_at),
                request.message,
            ]
        )
    for decision in report.case_decisions:
        sheet.append([decision_label(decision.decision), local_time(decision.decided_at), decision.reason or ""])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _autosize(sheet, [34, 18, 60])
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(HISTORY_HEADERS))}{max(sheet.max_row, 1)}"
    )
