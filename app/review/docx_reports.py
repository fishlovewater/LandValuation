"""Pure `.docx` renderer for an immutable Review risk report.

Body copy is business-facing Chinese. Structured evidence renders as readable
field/value rows, never raw JSON. Bucket names, object keys, and storage URLs
are never embedded.
"""

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.review.report_presentation import (
    CORRECTION_STATUS_LABELS,
    DOCUMENT_TYPE_LABELS,
    EVENT_LABELS,
    FINDING_TYPE_LABELS,
    OUTCOME_LABELS,
    REVIEW_STATUS_LABELS,
    RISK_LABELS,
    SEVERITY_LABELS,
    SOURCE_LABELS,
    TRIAGE_LABELS,
    URGENCY_LABELS,
    decision_label,
    district_label,
    label,
    local_time,
    readable_entries,
    skipped_rule_text,
)
from app.review.reports import ReviewReport


def _key_value_table(document: Document, rows: list[tuple[str, Any]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = "" if value is None else str(value)
        for paragraph in cells[0].paragraphs:
            for run in paragraph.runs:
                run.bold = True


def build_review_docx(report: ReviewReport) -> bytes:
    document = Document()
    document.core_properties.title = "土地估價審查報告"
    title = document.add_heading("土地估價審查報告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        document.styles[style_name].font.name = "Microsoft JhengHei"
    document.styles["Normal"].font.size = Pt(10.5)

    document.add_heading("案件基本資料", level=1)
    _key_value_table(
        document,
        [
            ("案件編號", report.case.case_no),
            ("案件名稱", report.case.case_title),
            ("行政區", district_label(report.case.district_code)),
            ("價格基準日", report.case.valuation_base_date),
            ("審查狀態", label(REVIEW_STATUS_LABELS, report.review_status)),
            ("審查完成時間", local_time(report.run.completed_at)),
        ],
    )

    document.add_heading("檢核覆蓋率", level=1)
    coverage = report.review_coverage
    if coverage is None:
        document.add_paragraph("此歷史批次未保存檢核覆蓋率資訊。")
    else:
        _key_value_table(
            document,
            [
                ("適用規則總數", coverage.total_rule_count),
                ("已執行規則", coverage.executed_rule_count),
                ("未執行規則", coverage.skipped_rule_count),
            ],
        )
        if coverage.skipped_rules:
            document.add_paragraph("未執行不代表通過；以下項目因資料不足未進行自動檢核。")
            for item in coverage.skipped_rules:
                document.add_paragraph(skipped_rule_text(item), style="List Bullet")

    provenance = report.input_provenance
    document.add_heading("審查文件", level=1)
    if provenance is None:
        document.add_paragraph("此報告未提供審查輸入版本資訊。")
    else:
        source_label = label(SOURCE_LABELS, provenance.source)
        _key_value_table(
            document,
            [
                ("資料來源", source_label),
                ("輸入版本", f"v{provenance.version_no}" if provenance.version_no is not None else ""),
                ("資料確認時間", local_time(provenance.frozen_at)),
            ],
        )
        if provenance.documents:
            document.add_paragraph("本次檢核使用文件", style="Intense Quote")
            for item in provenance.documents:
                filename = item.original_filename or "未命名文件"
                document.add_paragraph(
                    f"{filename}／{label(DOCUMENT_TYPE_LABELS, item.document_type)}／v{item.version_no}",
                    style="List Bullet",
                )

    document.add_heading("風險與期限", level=1)
    document.add_paragraph(
        "內容風險與期限緊急度為兩個獨立指標：風險反映報告內容問題，"
        "緊急度僅反映距離修正期限的天數。"
    )
    urgency = report.urgency
    _key_value_table(
        document,
        [
            ("內容風險等級", label(RISK_LABELS, report.risk_summary.overall_risk_level)),
            ("高風險疑點數", report.risk_summary.high_count),
            ("中風險疑點數", report.risk_summary.medium_count),
            ("低風險疑點數", report.risk_summary.low_count),
            ("缺件數", report.missing_item_count),
            (
                "期限緊急度",
                URGENCY_LABELS.get(urgency.level, urgency.level)
                if urgency
                else "未設定",
            ),
            ("剩餘天數", urgency.remaining_days if urgency else ""),
            ("修正期限", local_time(urgency.due_at) if urgency else ""),
        ],
    )

    document.add_heading("疑點與證據", level=1)
    if not report.findings:
        document.add_paragraph("本次檢核未發現疑點。")
    for finding in report.findings:
        document.add_heading(f"{finding.title}（{label(SEVERITY_LABELS, finding.severity)}風險）", level=2)
        status_label = TRIAGE_LABELS.get(finding.status, decision_label(finding.status))
        _key_value_table(
            document,
            [
                ("疑點類型", label(FINDING_TYPE_LABELS, finding.finding_type)),
                ("人工判定", status_label),
                ("疑點說明", f"{finding.title}：{finding.description}"),
                ("原報告內容", finding.reported_text or ""),
                ("原報告數值", finding.reported_value or ""),
            ],
        )
        document.add_paragraph("法規依據", style="Intense Quote")
        for line in readable_entries(finding.legal_basis) or ["（無）"]:
            document.add_paragraph(line, style="List Bullet")
        document.add_paragraph("查核證據", style="Intense Quote")
        for line in readable_entries(finding.source_evidence) or ["（無）"]:
            document.add_paragraph(line, style="List Bullet")
        if finding.ai_assessment.reasoning_summary:
            document.add_paragraph("AI 輔助說明", style="Intense Quote")
            document.add_paragraph(finding.ai_assessment.reasoning_summary)
        for decision in finding.decisions:
            document.add_paragraph(
                f"人工紀錄：{decision_label(decision.decision)}／{decision.reason}"
                f"（{local_time(decision.decided_at)}）"
            )

    document.add_heading("修正要求", level=1)
    if not report.correction_requests:
        document.add_paragraph("本批次未發出修正通知。")
    for request in report.correction_requests:
        document.add_heading(
            f"第 {request.request_no} 次修正通知（{label(CORRECTION_STATUS_LABELS, request.status)}）", level=2
        )
        _key_value_table(
            document,
            [
                ("通知內容", request.message),
                ("修正期限", local_time(request.due_at)),
                ("送出時間", local_time(request.sent_at)),
                ("原版本", request.base_document_version),
                ("回件版本", request.response_document_version or ""),
                ("回件時間", local_time(request.resubmitted_at)),
            ],
        )
        for item in request.items:
            document.add_paragraph(
                item.issue_summary, style="List Number"
            )
            document.add_paragraph(
                f"要求修正內容：{item.requested_correction}", style="List Bullet"
            )

    document.add_heading("新版重檢結果", level=1)
    rechecked_items = [
        (request, item)
        for request in report.correction_requests
        for item in request.items
    ]
    if not rechecked_items:
        document.add_paragraph("尚無新版重檢結果。")
    for request, item in rechecked_items:
        outcome = label(OUTCOME_LABELS, item.recheck_outcome)
        document.add_paragraph(
            f"第 {request.request_no} 次：{item.issue_summary}－{outcome}"
            f"（{local_time(request.rechecked_at)}）",
            style="List Bullet",
        )

    document.add_heading("審查結論", level=1)
    document.add_paragraph(f"審查狀態：{label(REVIEW_STATUS_LABELS, report.review_status)}")
    if report.review_status != "REVIEW_COMPLETED":
        document.add_paragraph("智慧審查已完成；本報告仍可由審查人員進一步人工確認與補充決策。")
    if not report.case_decisions and not report.history:
        document.add_paragraph("本批次尚無案件層級審查紀錄。")
    for event in report.history:
        document.add_paragraph(
            f"{label(EVENT_LABELS, event.event_type)}（{local_time(event.occurred_at)}）：{event.reason or ''}",
            style="List Bullet",
        )
    for decision in report.case_decisions:
        document.add_paragraph(
            f"{decision_label(decision.decision)}（{local_time(decision.decided_at)}）：{decision.reason or ''}",
            style="List Bullet",
        )

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
