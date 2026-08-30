from datetime import UTC, datetime, timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.review.reports import (
    ReportCorrectionItem,
    ReportCorrectionRequest,
    ReportHistoryEvent,
    ReportUrgency,
    build_review_report,
)
from app.review.tests.test_reports import report_input
from app.review.xlsx_reports import build_review_xlsx


@pytest.fixture
def report_fixture():
    data = report_input()
    finding_id = data.findings[0]["finding_id"]
    now = datetime.now(UTC)
    data.urgency = ReportUrgency(
        level="URGENT", remaining_days=2, due_at=now + timedelta(days=2)
    )
    data.correction_requests = [
        ReportCorrectionRequest(
            correction_request_id=uuid4(),
            request_no=1,
            status="RECHECKED",
            due_at=now + timedelta(days=2),
            message="請依附件疑點更正原報告",
            base_document_id=uuid4(),
            base_document_version=1,
            response_document_id=uuid4(),
            response_document_version=2,
            sent_at=now,
            resubmitted_at=now,
            rechecked_at=now,
            items=[
                ReportCorrectionItem(
                    finding_id=finding_id,
                    finding_code="RATE-001",
                    severity="HIGH",
                    page_number=3,
                    reported_text="-12%",
                    reported_value="-12",
                    legal_basis=[{"article": "第10條"}],
                    source_evidence=[{"excerpt": "Adjustment rate -12%"}],
                    issue_summary="調整率不一致",
                    requested_correction="請更正調整率並補充依據",
                    recheck_outcome="RESOLVED",
                    resulting_finding_id=None,
                )
            ],
        )
    ]
    data.history = [
        ReportHistoryEvent(
            event_type="CORRECTION_SENT",
            occurred_at=now,
            actor_id=uuid4(),
            reason="送出第 1 次修正通知",
        )
    ]
    return build_review_report(data)


def test_xlsx_has_required_sheets_and_no_macros(report_fixture):
    content = build_review_xlsx(report_fixture)
    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames == [
        "案件摘要",
        "疑點與修正要求",
        "新版重檢結果",
        "審查歷程",
    ]
    assert workbook.vba_archive is None


def test_xlsx_finding_rows_contain_evidence_not_formal_value(report_fixture):
    workbook = load_workbook(BytesIO(build_review_xlsx(report_fixture)))
    headers = [cell.value for cell in workbook["疑點與修正要求"][1]]
    assert "原報告頁碼" in headers
    assert "法規依據" in headers
    assert "建議修正方向" in headers
    assert "正式採用值" not in headers


def test_xlsx_reopens_and_contains_case_and_urgency(report_fixture):
    workbook = load_workbook(BytesIO(build_review_xlsx(report_fixture)))
    summary_text = "\n".join(
        str(cell.value)
        for row in workbook["案件摘要"].iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert report_fixture.case.case_no in summary_text
    # Deadline urgency is reported separately from content risk.
    assert "期限緊急度" in summary_text
    assert "內容風險等級" in summary_text


def test_xlsx_recheck_sheet_lists_outcomes(report_fixture):
    workbook = load_workbook(BytesIO(build_review_xlsx(report_fixture)))
    sheet = workbook["新版重檢結果"]
    values = [
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    assert "RESOLVED" in values


def test_xlsx_exposes_no_storage_internals(report_fixture):
    content = build_review_xlsx(report_fixture)
    workbook = load_workbook(BytesIO(content))
    text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "localhost" not in text
    assert "land-valuation" not in text
    assert "object_key" not in text
