from io import BytesIO

from docx import Document

from app.review.docx_reports import build_review_docx
from app.review.tests.test_xlsx_reports import report_fixture  # noqa: F401


def _all_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_docx_contains_required_sections(report_fixture):  # noqa: F811
    document = Document(BytesIO(build_review_docx(report_fixture)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for heading in [
        "案件基本資料",
        "檢核覆蓋率",
        "審查文件",
        "風險與期限",
        "疑點與證據",
        "修正要求",
        "新版重檢結果",
        "審查結論",
    ]:
        assert heading in text
    assert "審查人員另訂正式值" not in text


def test_docx_contains_review_input_provenance(report_fixture):  # noqa: F811
    document = Document(BytesIO(build_review_docx(report_fixture)))
    text = _all_text(document)
    assert "外部案件" in text
    assert "external-report-v2.pdf" in text
    assert "SHA-256" not in text
    assert "a" * 64 not in text


def test_docx_labels_ai_output_and_keeps_sources(report_fixture):  # noqa: F811
    document = Document(BytesIO(build_review_docx(report_fixture)))
    text = _all_text(document)
    assert "AI 輔助說明" in text
    # Evidence excerpts must survive as readable rows.
    assert "Adjustment rate -12%" in text
    assert "條文：第10條" in text
    assert "article" not in text


def test_docx_marks_legacy_decision_as_history(report_fixture):  # noqa: F811
    # The fixture's stored decision is a legacy PARTIALLY_ACCEPTED record.
    document = Document(BytesIO(build_review_docx(report_fixture)))
    text = _all_text(document)
    assert "舊流程歷史決策" in text


def test_docx_separates_risk_from_deadline_urgency(report_fixture):  # noqa: F811
    document = Document(BytesIO(build_review_docx(report_fixture)))
    text = _all_text(document)
    assert "內容風險等級" in text
    assert "期限緊急度" in text


def test_docx_marks_skipped_checks_as_not_passed(report_fixture):  # noqa: F811
    document = Document(BytesIO(build_review_docx(report_fixture)))
    text = _all_text(document)
    assert "未執行不代表通過" in text
    assert "土地登記資料交叉檢核" in text
    assert "LAND_REGISTER_CROSSCHECK" not in text
    assert "LAND_REGISTER_AREA" not in text


def test_docx_exposes_no_storage_internals(report_fixture):  # noqa: F811
    document = Document(BytesIO(build_review_docx(report_fixture)))
    text = _all_text(document)
    assert "localhost" not in text
    assert "land-valuation" not in text
    assert "object_key" not in text
    assert "平台送審識別碼" not in text
    assert "快照格式" not in text
    assert "RATE-001" not in text
    assert "RATE_OUT_OF_RANGE" not in text
    assert "PARTIALLY_ACCEPTED" not in text
    assert "REVIEW_REQUIRED" not in text
