from __future__ import annotations

from io import BytesIO

from docx import Document
from fastapi.concurrency import run_in_threadpool
from pydantic import BaseModel


class DocumentTextPreviewTooLargeError(ValueError):
    pass


class DocumentTextPreview(BaseModel):
    kind: str = "text"
    text: str
    truncated: bool


def build_docx_text_preview(
    content: bytes,
    *,
    max_characters: int = 50_000,
) -> DocumentTextPreview:
    document = Document(BytesIO(content))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        table_rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                table_rows.append(" | ".join(cells))
        if table_rows:
            parts.append("[表格]\n" + "\n".join(table_rows))

    text = "\n\n".join(parts)
    truncated = len(text) > max_characters
    if truncated:
        text = text[:max_characters].rstrip()
    return DocumentTextPreview(text=text, truncated=truncated)


async def preview_storage_docx(
    storage,
    object_key: str,
    *,
    max_bytes: int = 10 * 1024 * 1024,
    max_characters: int = 50_000,
) -> DocumentTextPreview:
    response = await storage.download(object_key)
    try:
        content = await run_in_threadpool(response.read, max_bytes + 1)
    finally:
        response.close()
        response.release_conn()

    if len(content) > max_bytes:
        raise DocumentTextPreviewTooLargeError(
            f"document preview exceeds {max_bytes} bytes"
        )
    return await run_in_threadpool(
        build_docx_text_preview,
        content,
        max_characters=max_characters,
    )