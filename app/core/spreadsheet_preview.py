from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from fastapi.concurrency import run_in_threadpool
from openpyxl import load_workbook
from pydantic import BaseModel

XLS_MIME_TYPE = "application/vnd.ms-excel"
XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class SpreadsheetPreviewTooLargeError(ValueError):
    pass


class SpreadsheetPreviewSheet(BaseModel):
    name: str
    rows: list[list[str | int | float | bool | None]]
    total_rows: int
    total_columns: int
    truncated: bool


class SpreadsheetPreview(BaseModel):
    kind: str = "spreadsheet"
    sheets: list[SpreadsheetPreviewSheet]
    truncated: bool


def _cell_value(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _build_xlsx_preview(
    content: bytes,
    *,
    max_sheets: int = 8,
    max_rows: int = 80,
    max_columns: int = 24,
) -> SpreadsheetPreview:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheets: list[SpreadsheetPreviewSheet] = []
        workbook_truncated = len(workbook.sheetnames) > max_sheets
        for worksheet in workbook.worksheets[:max_sheets]:
            total_rows = worksheet.max_row or 0
            total_columns = worksheet.max_column or 0
            rows: list[list[str | int | float | bool | None]] = []
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=min(total_rows, max_rows),
                min_col=1,
                max_col=min(total_columns, max_columns),
                values_only=True,
            ):
                rows.append([_cell_value(value) for value in row])
            truncated = total_rows > max_rows or total_columns > max_columns
            workbook_truncated = workbook_truncated or truncated
            sheets.append(
                SpreadsheetPreviewSheet(
                    name=worksheet.title,
                    rows=rows,
                    total_rows=total_rows,
                    total_columns=total_columns,
                    truncated=truncated,
                )
            )
        return SpreadsheetPreview(sheets=sheets, truncated=workbook_truncated)
    finally:
        workbook.close()


def _build_xls_preview(
    content: bytes,
    *,
    max_sheets: int = 8,
    max_rows: int = 80,
    max_columns: int = 24,
) -> SpreadsheetPreview:
    # xlrd 2.x intentionally supports only the legacy binary .xls format.
    # Import lazily so normal XLSX previews do not depend on it at import time.
    import xlrd

    workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        sheets: list[SpreadsheetPreviewSheet] = []
        workbook_truncated = workbook.nsheets > max_sheets
        for worksheet in workbook.sheets()[:max_sheets]:
            total_rows = worksheet.nrows
            total_columns = worksheet.ncols
            rows: list[list[str | int | float | bool | None]] = []
            for row_index in range(min(total_rows, max_rows)):
                row: list[str | int | float | bool | None] = []
                for column_index in range(min(total_columns, max_columns)):
                    cell = worksheet.cell(row_index, column_index)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        value: str | int | float | bool | None = None
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(cell.value)
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(cell.value, workbook.datemode).isoformat()
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        number = float(cell.value)
                        value = int(number) if number.is_integer() else number
                    else:
                        value = _cell_value(cell.value)
                    row.append(value)
                rows.append(row)
            truncated = total_rows > max_rows or total_columns > max_columns
            workbook_truncated = workbook_truncated or truncated
            sheets.append(
                SpreadsheetPreviewSheet(
                    name=worksheet.name,
                    rows=rows,
                    total_rows=total_rows,
                    total_columns=total_columns,
                    truncated=truncated,
                )
            )
        return SpreadsheetPreview(sheets=sheets, truncated=workbook_truncated)
    finally:
        workbook.release_resources()


def build_spreadsheet_preview(
    content: bytes,
    *,
    mime_type: str = XLSX_MIME_TYPE,
    max_sheets: int = 8,
    max_rows: int = 80,
    max_columns: int = 24,
) -> SpreadsheetPreview:
    if mime_type == XLS_MIME_TYPE:
        return _build_xls_preview(
            content,
            max_sheets=max_sheets,
            max_rows=max_rows,
            max_columns=max_columns,
        )
    if mime_type != XLSX_MIME_TYPE:
        raise ValueError(f"unsupported spreadsheet MIME type: {mime_type}")
    return _build_xlsx_preview(
        content,
        max_sheets=max_sheets,
        max_rows=max_rows,
        max_columns=max_columns,
    )


async def preview_storage_xlsx(
    storage,
    object_key: str,
    *,
    max_bytes: int = 10 * 1024 * 1024,
) -> SpreadsheetPreview:
    response = await storage.download(object_key)
    try:
        content = await run_in_threadpool(response.read, max_bytes + 1)
    finally:
        response.close()
        response.release_conn()
    if len(content) > max_bytes:
        raise SpreadsheetPreviewTooLargeError(
            f"spreadsheet preview exceeds {max_bytes} bytes"
        )
    return await run_in_threadpool(build_spreadsheet_preview, content)


async def preview_storage_spreadsheet(
    storage,
    object_key: str,
    *,
    mime_type: str,
    max_bytes: int = 10 * 1024 * 1024,
) -> SpreadsheetPreview:
    response = await storage.download(object_key)
    try:
        content = await run_in_threadpool(response.read, max_bytes + 1)
    finally:
        response.close()
        response.release_conn()
    if len(content) > max_bytes:
        raise SpreadsheetPreviewTooLargeError(
            f"spreadsheet preview exceeds {max_bytes} bytes"
        )
    return await run_in_threadpool(
        build_spreadsheet_preview,
        content,
        mime_type=mime_type,
    )
