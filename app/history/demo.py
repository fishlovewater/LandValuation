"""Deterministic development data for manually testing Case-history.

Run inside the API runtime with ``python -m app.history.demo seed|status|reset``.
"""

import argparse
import asyncio
import hashlib
import os
import sys
from io import BytesIO
from pathlib import Path
from uuid import UUID

from sqlalchemy import bindparam, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine


def _load_dotenv() -> None:
    path = Path.cwd() / ".env"
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


_load_dotenv()
if os.getenv("POSTGRES_HOST"):
    # Compose explicitly supplies POSTGRES_HOST=db. Prefer these component
    # variables over a mounted .env DATABASE_URL that points at localhost.
    DATABASE_URL = (
        "postgresql+psycopg://"
        f"{os.getenv('POSTGRES_USER', 'app_user')}:"
        f"{os.getenv('POSTGRES_PASSWORD', 'change_me')}@"
        f"{os.getenv('POSTGRES_HOST')}:"
        f"{os.getenv('POSTGRES_PORT', '5432')}/"
        f"{os.getenv('POSTGRES_DB', 'land_valuation')}"
    )
else:
    DATABASE_URL = os.getenv("DATABASE_URL") or (
        "postgresql+psycopg://app_user:change_me@localhost:5432/land_valuation"
    )
engine = None
SessionFactory = None


def _database_runtime():
    global engine, SessionFactory
    if engine is not None:
        return engine, SessionFactory
    try:
        engine = create_async_engine(DATABASE_URL, pool_pre_ping=True)
    except (ImportError, ModuleNotFoundError) as exc:
        raise RuntimeError(
            "History demo requires a working psycopg driver (for example "
            "psycopg[binary]); run it in the complete API runtime/container"
        ) from exc
    SessionFactory = async_sessionmaker(engine, expire_on_commit=False)
    return engine, SessionFactory


def _minio_client():
    try:
        from minio import Minio
    except ImportError as exc:
        raise RuntimeError(
            "seed requires the project's 'minio' package; install project "
            "requirements or run this command inside the API container"
        ) from exc
    access_key = os.getenv("MINIO_ACCESS_KEY") or os.getenv("MINIO_ROOT_USER")
    secret_key = os.getenv("MINIO_SECRET_KEY") or os.getenv("MINIO_ROOT_PASSWORD")
    if not access_key or not secret_key:
        raise RuntimeError("MINIO access key/secret are missing from environment or .env")
    return Minio(
        os.getenv("MINIO_ENDPOINT", "localhost:9000"),
        access_key=access_key,
        secret_key=secret_key,
        secure=os.getenv("MINIO_SECURE", "false").lower() == "true",
    )


def _bucket() -> str:
    return os.getenv("MINIO_BUCKET", "land-valuation")


def _ensure_development() -> None:
    if os.getenv("APP_ENV", "").lower() != "development":
        raise RuntimeError(
            "DEVELOPMENT_ONLY: demo commands require APP_ENV=development"
        )


VALUATION_CASE_ID = UUID("71000000-0000-4000-8000-000000000001")
REVIEW_CASE_ID = UUID("71000000-0000-4000-8000-000000000002")
BOTH_CASE_ID = UUID("71000000-0000-4000-8000-000000000003")
CASE_IDS = (VALUATION_CASE_ID, REVIEW_CASE_ID, BOTH_CASE_ID)
PARCEL_ID = UUID("72000000-0000-4000-8000-000000000001")
FORM_ID = UUID("73000000-0000-4000-8000-000000000001")
BOTH_FORM_ID = UUID("73000000-0000-4000-8000-000000000002")
FORM_IDS = (FORM_ID, BOTH_FORM_ID)
VALUATION_ID = UUID("74000000-0000-4000-8000-000000000001")
REVIEW_ID = UUID("75000000-0000-4000-8000-000000000001")
BOTH_REVIEW_ID = UUID("75000000-0000-4000-8000-000000000002")
REVIEW_IDS = (REVIEW_ID, BOTH_REVIEW_ID)
RISK_ID = UUID("76000000-0000-4000-8000-000000000001")
DOWNLOAD_DOCUMENT_ID = UUID("77000000-0000-4000-8000-000000000001")
MISSING_DOCUMENT_ID = UUID("77000000-0000-4000-8000-000000000002")
DOCX_DOCUMENT_ID = UUID("77000000-0000-4000-8000-000000000003")
XLSX_V1_DOCUMENT_ID = UUID("77000000-0000-4000-8000-000000000004")
XLSX_V2_DOCUMENT_ID = UUID("77000000-0000-4000-8000-000000000005")
XLSX_DOCUMENT_GROUP_ID = UUID("77000000-0000-4000-8000-000000000100")
DOCUMENT_IDS = (
    DOWNLOAD_DOCUMENT_ID,
    MISSING_DOCUMENT_ID,
    DOCX_DOCUMENT_ID,
    XLSX_V1_DOCUMENT_ID,
    XLSX_V2_DOCUMENT_ID,
)
DOWNLOAD_OBJECT_KEY = f"cases/{BOTH_CASE_ID}/generated/{DOWNLOAD_DOCUMENT_ID}/v1/history-demo-report.pdf"
MISSING_OBJECT_KEY = f"cases/{REVIEW_CASE_ID}/generated/{MISSING_DOCUMENT_ID}/v1/history-demo-missing.docx"
DOCX_OBJECT_KEY = f"cases/{BOTH_CASE_ID}/generated/{DOCX_DOCUMENT_ID}/v1/history-demo-notes.docx"
XLSX_V1_OBJECT_KEY = f"cases/{BOTH_CASE_ID}/generated/{XLSX_DOCUMENT_GROUP_ID}/v1/history-demo-values-v1.xlsx"
XLSX_V2_OBJECT_KEY = f"cases/{BOTH_CASE_ID}/generated/{XLSX_DOCUMENT_GROUP_ID}/v2/history-demo-values-v2.xlsx"

EXTRACTION_V1_ID = UUID("79000000-0000-4000-8000-000000000001")
EXTRACTION_V2_ID = UUID("79000000-0000-4000-8000-000000000002")
EXTRACTION_IDS = (EXTRACTION_V1_ID, EXTRACTION_V2_ID)
EXTRACTED_FIELD_V1_ID = UUID("7a000000-0000-4000-8000-000000000001")
EXTRACTED_FIELD_V2_ID = UUID("7a000000-0000-4000-8000-000000000002")
EXTRACTED_FIELD_IDS = (EXTRACTED_FIELD_V1_ID, EXTRACTED_FIELD_V2_ID)
CASE_VERSION_V1_ID = UUID("7b000000-0000-4000-8000-000000000001")
CASE_VERSION_V2_ID = UUID("7b000000-0000-4000-8000-000000000002")
CHANGE_LOG_ID = UUID("7c000000-0000-4000-8000-000000000001")

APPRAISER_USER_ID = UUID("78000000-0000-4000-8000-000000000001")
REVIEWER_USER_ID = UUID("78000000-0000-4000-8000-000000000002")
USER_IDS = (APPRAISER_USER_ID, REVIEWER_USER_ID)
DEMO_PASSWORD = "HistoryDemo123!"
DEMO_USERS = (
    {
        "user_id": APPRAISER_USER_ID,
        "username": "history_appraiser",
        "email": "history_appraiser@example.test",
        "display_name": "History 測試估價人員",
        "role_code": "APPRAISER",
    },
    {
        "user_id": REVIEWER_USER_ID,
        "username": "history_reviewer",
        "email": "history_reviewer@example.test",
        "display_name": "History 測試審查人員",
        "role_code": "REVIEWER",
    },
)

PDF_BYTES = b"""%PDF-1.4
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 144]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj
4 0 obj<</Length 57>>stream
BT /F1 16 Tf 30 80 Td (Case-history demo download) Tj ET
endstream endobj
5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
xref
0 6
0000000000 65535 f\x20
0000000009 00000 n\x20
0000000058 00000 n\x20
0000000115 00000 n\x20
0000000241 00000 n\x20
0000000347 00000 n\x20
trailer<</Size 6/Root 1 0 R>>
startxref
417
%%EOF
"""


def _build_docx_bytes() -> bytes:
    from docx import Document

    buffer = BytesIO()
    document = Document()
    document.add_heading("案件歷史 DOCX 預覽測試", level=1)
    document.add_paragraph("此文件用於驗證案件歷史的 Word 文字預覽功能。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "欄位"
    table.cell(0, 1).text = "內容"
    table.cell(1, 0).text = "案件編號"
    table.cell(1, 1).text = "HIST-BOTH-001"
    document.save(buffer)
    return buffer.getvalue()


def _build_xlsx_bytes(version: int, comparison_price: int) -> bytes:
    from openpyxl import Workbook

    buffer = BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "比準地估價"
    worksheet.append(["文件版本", "比較法價格", "案件編號"])
    worksheet.append([version, comparison_price, "HIST-BOTH-001"])
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


UPLOAD_OBJECT_KEYS = (
    DOWNLOAD_OBJECT_KEY,
    DOCX_OBJECT_KEY,
    XLSX_V1_OBJECT_KEY,
    XLSX_V2_OBJECT_KEY,
)


async def _delete_rows(session) -> None:
    await session.execute(text("DELETE FROM review.risk_summaries WHERE review_id=:id"), {"id": REVIEW_ID})
    await session.execute(
        text("DELETE FROM review.reviews WHERE review_id IN :ids").bindparams(
            bindparam("ids", expanding=True)
        ),
        {"ids": REVIEW_IDS},
    )
    await session.execute(
        text("DELETE FROM history.change_logs WHERE case_id IN :ids").bindparams(
            bindparam("ids", expanding=True)
        ),
        {"ids": CASE_IDS},
    )
    await session.execute(
        text("DELETE FROM history.case_versions WHERE case_id IN :ids").bindparams(
            bindparam("ids", expanding=True)
        ),
        {"ids": CASE_IDS},
    )
    await session.execute(
        text("DELETE FROM valuation.extracted_fields WHERE extracted_field_id IN :ids").bindparams(
            bindparam("ids", expanding=True)
        ),
        {"ids": EXTRACTED_FIELD_IDS},
    )
    await session.execute(
        text("DELETE FROM valuation.document_extractions WHERE extraction_id IN :ids").bindparams(
            bindparam("ids", expanding=True)
        ),
        {"ids": EXTRACTION_IDS},
    )
    await session.execute(text("DELETE FROM valuation.valuations WHERE valuation_id=:id"), {"id": VALUATION_ID})
    await session.execute(
        text("DELETE FROM valuation.form_instances WHERE form_instance_id IN :ids").bindparams(
            bindparam("ids", expanding=True)
        ),
        {"ids": FORM_IDS},
    )
    await session.execute(text("DELETE FROM valuation.parcels WHERE parcel_id=:id"), {"id": PARCEL_ID})
    await session.execute(
        text("DELETE FROM valuation.documents WHERE document_id IN :ids").bindparams(bindparam("ids", expanding=True)),
        {"ids": DOCUMENT_IDS},
    )
    await session.execute(
        text("DELETE FROM valuation.cases WHERE case_id IN :ids").bindparams(bindparam("ids", expanding=True)),
        {"ids": CASE_IDS},
    )
    await session.execute(
        text("DELETE FROM auth.user_roles WHERE user_id IN :ids").bindparams(
            bindparam("ids", expanding=True)
        ),
        {"ids": USER_IDS},
    )
    await session.execute(
        text("DELETE FROM auth.users WHERE user_id IN :ids").bindparams(
            bindparam("ids", expanding=True)
        ),
        {"ids": USER_IDS},
    )


async def reset() -> None:
    _ensure_development()
    _, sessions = _database_runtime()
    async with sessions() as session:
        async with session.begin():
            await _delete_rows(session)
    client = _minio_client()
    for object_key in UPLOAD_OBJECT_KEYS:
        await asyncio.to_thread(client.remove_object, _bucket(), object_key)


async def seed() -> None:
    _ensure_development()
    # Lazy import keeps ``status``/``reset`` usable in a partial local Python
    # environment; password hashing is required only while creating accounts.
    from app.core.security import hash_password

    client = _minio_client()  # Validate dependency and credentials before mutation.
    bucket = _bucket()
    if not await asyncio.to_thread(client.bucket_exists, bucket):
        raise RuntimeError(f"MinIO bucket does not exist: {bucket}")
    await reset()
    docx_bytes = _build_docx_bytes()
    xlsx_v1_bytes = _build_xlsx_bytes(1, 120000)
    xlsx_v2_bytes = _build_xlsx_bytes(2, 125000)
    upload_objects = (
        (DOWNLOAD_OBJECT_KEY, PDF_BYTES, "application/pdf"),
        (DOCX_OBJECT_KEY, docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        (XLSX_V1_OBJECT_KEY, xlsx_v1_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        (XLSX_V2_OBJECT_KEY, xlsx_v2_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    )
    uploaded_etags: dict[str, str] = {}
    for object_key, payload, content_type in upload_objects:
        uploaded = await asyncio.to_thread(
            client.put_object,
            bucket,
            object_key,
            BytesIO(payload),
            len(payload),
            content_type=content_type,
        )
        uploaded_etags[object_key] = uploaded.etag
    try:
        _, sessions = _database_runtime()
        async with sessions() as session:
            async with session.begin():
                for demo_user in DEMO_USERS:
                    role_exists = await session.scalar(
                        text("""SELECT EXISTS (
                            SELECT 1 FROM auth.roles
                            WHERE role_code=:role_code AND is_active=true
                        )"""),
                        {"role_code": demo_user["role_code"]},
                    )
                    if not role_exists:
                        raise RuntimeError(
                            f"{demo_user['role_code']}_ROLE_MISSING: run migrations before seed"
                        )
                    await session.execute(text("""INSERT INTO auth.users
                        (user_id,username,email,password_hash,display_name,is_active,created_at,updated_at)
                        VALUES (:user_id,:username,:email,:password_hash,:display_name,true,now(),now())"""), {
                        **demo_user,
                        "password_hash": hash_password(DEMO_PASSWORD),
                    })
                    await session.execute(text("""INSERT INTO auth.user_roles (user_id,role_id)
                        SELECT :user_id,role_id FROM auth.roles
                        WHERE role_code=:role_code AND is_active=true"""), demo_user)
                await session.execute(text("""INSERT INTO valuation.cases
                    (case_id,case_no,case_title,case_type,valuation_base_date,city_code,district_code,land_use_type,case_status)
                    VALUES
                    (:v,'HIST-VAL-001','History 測試－只有估價結構化資料','LAND',DATE '2026-08-01','31','3101','COMMERCIAL','CORRECTION'),
                    (:r,'HIST-REV-001','History 測試－Review metadata 但 MinIO 缺檔','LAND',DATE '2026-08-02','31','3102','RESIDENTIAL','REVIEWING'),
                    (:b,'HIST-BOTH-001','History 測試－雙子系統及可下載文件','LAND',DATE '2026-08-03','31','3104','COMMERCIAL','COMPLETED')"""),
                    {"v": VALUATION_CASE_ID, "r": REVIEW_CASE_ID, "b": BOTH_CASE_ID})
                await session.execute(text("""INSERT INTO valuation.parcels
                    (parcel_id,case_id,district_code,section_name,subsection_name,land_no,area_sqm)
                    VALUES (:id,:case_id,'3101','測試段','','123-4',168.5)"""),
                    {"id": PARCEL_ID, "case_id": VALUATION_CASE_ID})
                await session.execute(text("""INSERT INTO valuation.form_instances
                    (form_instance_id,case_id,form_code,version_no,form_status)
                    VALUES
                    (:id,:case_id,'F01',1,'FINAL'),
                    (:both_id,:both_case,'F03',1,'FINAL')"""),
                    {
                        "id": FORM_ID,
                        "case_id": VALUATION_CASE_ID,
                        "both_id": BOTH_FORM_ID,
                        "both_case": BOTH_CASE_ID,
                    })
                await session.execute(text("""INSERT INTO valuation.valuations
                    (valuation_id,case_id,form_instance_id,valuation_type,unit_price,total_value,calculation_snapshot,result_status)
                    VALUES (:id,:case_id,:form_id,'CASE',88000,14828000,CAST(:snapshot AS jsonb),'FINAL')"""),
                    {"id": VALUATION_ID, "case_id": VALUATION_CASE_ID, "form_id": FORM_ID,
                     "snapshot": '{"demo":true,"method":"comparison","unit_price":88000}'})
                await session.execute(text("""INSERT INTO review.reviews
                    (review_id,case_id,review_type,review_status,received_at,started_at,completed_at)
                    VALUES (:id,:case_id,'SMART_REVIEW','REVIEW_COMPLETED',
                    TIMESTAMPTZ '2026-08-19 09:00:00+08',
                    TIMESTAMPTZ '2026-08-20 10:00:00+08',
                    TIMESTAMPTZ '2026-08-20 12:00:00+08'),
                    (:both_id,:both_case,'MANUAL_REVIEW','REVIEW_COMPLETED',
                    TIMESTAMPTZ '2026-08-20 09:00:00+08',
                    TIMESTAMPTZ '2026-08-21 10:00:00+08',
                    TIMESTAMPTZ '2026-08-21 12:00:00+08')"""),
                    {"id": REVIEW_ID, "case_id": REVIEW_CASE_ID,
                     "both_id": BOTH_REVIEW_ID, "both_case": BOTH_CASE_ID})
                await session.execute(text("""INSERT INTO review.risk_summaries
                    (risk_summary_id,review_id,overall_risk_level,risk_score,summary,category_scores)
                    VALUES (:id,:review_id,'LOW',12.5,'History demo：低風險測試資料',
                    CAST(:category_scores AS jsonb))"""),
                    {"id": RISK_ID, "review_id": REVIEW_ID,
                     "category_scores": '{"completeness":95}'})
                await session.execute(text("""INSERT INTO valuation.documents
                    (document_id,document_group_id,case_id,document_type,original_filename,mime_type,
                     bucket_name,object_key,checksum_sha256,file_size_bytes,version_no,is_active,storage_etag)
                    VALUES
                    (:did,:did,:both,'review-report','history-demo-report.pdf','application/pdf',
                     :bucket,:dkey,:dhash,:dsize,1,true,:detag),
                    (:mid,:mid,:review,'review-report','history-demo-missing.docx',
                     'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                     :bucket,:mkey,:mhash,128,1,true,NULL),
                    (:docx,:docx,:both,'review-report','history-demo-notes.docx',
                     'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                     :bucket,:docxkey,:docxhash,:docxsize,1,true,:docxetag),
                    (:x1,:xgroup,:both,'generated-report','history-demo-values-v1.xlsx',
                     'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     :bucket,:x1key,:x1hash,:x1size,1,false,:x1etag),
                    (:x2,:xgroup,:both,'generated-report','history-demo-values-v2.xlsx',
                     'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     :bucket,:x2key,:x2hash,:x2size,2,true,:x2etag)"""), {
                    "did": DOWNLOAD_DOCUMENT_ID, "both": BOTH_CASE_ID,
                    "mid": MISSING_DOCUMENT_ID, "review": REVIEW_CASE_ID,
                    "bucket": bucket, "dkey": DOWNLOAD_OBJECT_KEY,
                    "dhash": hashlib.sha256(PDF_BYTES).hexdigest(), "dsize": len(PDF_BYTES),
                    "detag": uploaded_etags[DOWNLOAD_OBJECT_KEY], "mkey": MISSING_OBJECT_KEY,
                    "mhash": hashlib.sha256(b"missing-demo-object").hexdigest(),
                    "docx": DOCX_DOCUMENT_ID,
                    "docxkey": DOCX_OBJECT_KEY,
                    "docxhash": hashlib.sha256(docx_bytes).hexdigest(),
                    "docxsize": len(docx_bytes),
                    "docxetag": uploaded_etags[DOCX_OBJECT_KEY],
                    "x1": XLSX_V1_DOCUMENT_ID,
                    "x2": XLSX_V2_DOCUMENT_ID,
                    "xgroup": XLSX_DOCUMENT_GROUP_ID,
                    "x1key": XLSX_V1_OBJECT_KEY,
                    "x2key": XLSX_V2_OBJECT_KEY,
                    "x1hash": hashlib.sha256(xlsx_v1_bytes).hexdigest(),
                    "x2hash": hashlib.sha256(xlsx_v2_bytes).hexdigest(),
                    "x1size": len(xlsx_v1_bytes),
                    "x2size": len(xlsx_v2_bytes),
                    "x1etag": uploaded_etags[XLSX_V1_OBJECT_KEY],
                    "x2etag": uploaded_etags[XLSX_V2_OBJECT_KEY],
                })
                await session.execute(text("""INSERT INTO valuation.document_extractions
                    (extraction_id,case_id,document_id,provider,extraction_status,extracted_text,
                     page_count,created_by_user_id,started_at,completed_at)
                    VALUES
                    (:e1,:case_id,:x1,'LOCAL_XLSX','COMPLETED','比較法價格 120000',1,:user_id,
                     TIMESTAMPTZ '2026-08-22 09:00:00+08',TIMESTAMPTZ '2026-08-22 09:01:00+08'),
                    (:e2,:case_id,:x2,'LOCAL_XLSX','COMPLETED','比較法價格 125000',1,:user_id,
                     TIMESTAMPTZ '2026-08-23 09:00:00+08',TIMESTAMPTZ '2026-08-23 09:01:00+08')"""), {
                    "e1": EXTRACTION_V1_ID,
                    "e2": EXTRACTION_V2_ID,
                    "case_id": BOTH_CASE_ID,
                    "x1": XLSX_V1_DOCUMENT_ID,
                    "x2": XLSX_V2_DOCUMENT_ID,
                    "user_id": APPRAISER_USER_ID,
                })
                await session.execute(text("""INSERT INTO valuation.extracted_fields
                    (extracted_field_id,case_id,extraction_id,document_id,form_code,field_name,
                     extracted_value,confidence,source_page,source_text,analysis_provider,field_status,
                     confirmed_value,confirmed_by_user_id,confirmed_at,applied_form_instance_id,applied_at)
                    VALUES
                    (:f1,:case_id,:e1,:x1,'F03','comparison_price',CAST('120000' AS jsonb),0.99,1,
                     '比較法價格 120000','RULE','APPLIED',CAST('120000' AS jsonb),:user_id,
                     TIMESTAMPTZ '2026-08-22 09:02:00+08',:form_id,TIMESTAMPTZ '2026-08-22 09:03:00+08'),
                    (:f2,:case_id,:e2,:x2,'F03','comparison_price',CAST('125000' AS jsonb),0.99,1,
                     '比較法價格 125000','RULE','APPLIED',CAST('125000' AS jsonb),:user_id,
                     TIMESTAMPTZ '2026-08-23 09:02:00+08',:form_id,TIMESTAMPTZ '2026-08-23 09:03:00+08')"""), {
                    "f1": EXTRACTED_FIELD_V1_ID,
                    "f2": EXTRACTED_FIELD_V2_ID,
                    "case_id": BOTH_CASE_ID,
                    "e1": EXTRACTION_V1_ID,
                    "e2": EXTRACTION_V2_ID,
                    "x1": XLSX_V1_DOCUMENT_ID,
                    "x2": XLSX_V2_DOCUMENT_ID,
                    "user_id": APPRAISER_USER_ID,
                    "form_id": BOTH_FORM_ID,
                })
                await session.execute(text("""INSERT INTO history.case_versions
                    (case_version_id,case_id,version_no,snapshot,change_summary,created_by_user_id,created_at)
                    VALUES
                    (:v1,:case_id,1,CAST(:s1 AS jsonb),'建立案件歷史驗證基準',:user_id,
                     TIMESTAMPTZ '2026-08-22 09:10:00+08'),
                    (:v2,:case_id,2,CAST(:s2 AS jsonb),'更新比準地價格並完成審查',:user_id,
                     TIMESTAMPTZ '2026-08-23 09:10:00+08')"""), {
                    "v1": CASE_VERSION_V1_ID,
                    "v2": CASE_VERSION_V2_ID,
                    "case_id": BOTH_CASE_ID,
                    "s1": '{"case_status":"REVIEWING","comparison_price":120000}',
                    "s2": '{"case_status":"COMPLETED","comparison_price":125000}',
                    "user_id": APPRAISER_USER_ID,
                })
                await session.execute(text("""INSERT INTO history.change_logs
                    (change_log_id,case_id,entity_type,entity_id,field_name,old_value,new_value,
                     change_reason,changed_by_user_id,changed_at)
                    VALUES (:id,:case_id,'case',:case_id,'case_status',CAST(:old AS jsonb),CAST(:new AS jsonb),
                     '完成案件歷史驗證流程',:user_id,TIMESTAMPTZ '2026-08-23 09:11:00+08')"""), {
                    "id": CHANGE_LOG_ID,
                    "case_id": BOTH_CASE_ID,
                    "old": '"REVIEWING"',
                    "new": '"COMPLETED"',
                    "user_id": APPRAISER_USER_ID,
                })
    except Exception:
        for object_key, _, _ in upload_objects:
            await asyncio.to_thread(client.remove_object, bucket, object_key)
        raise


async def status() -> None:
    _, sessions = _database_runtime()
    async with sessions() as session:
        rows = (await session.execute(
            text("""SELECT c.case_id,c.case_no,c.case_title,count(d.document_id) document_count
                    FROM valuation.cases c LEFT JOIN valuation.documents d ON d.case_id=c.case_id
                    WHERE c.case_id IN :ids GROUP BY c.case_id,c.case_no,c.case_title ORDER BY c.case_no""").bindparams(bindparam("ids", expanding=True)),
            {"ids": CASE_IDS})).mappings().all()
        users = (await session.execute(
            text("""SELECT u.username,u.display_name,r.role_code,u.is_active
                    FROM auth.users u
                    JOIN auth.user_roles ur ON ur.user_id=u.user_id
                    JOIN auth.roles r ON r.role_id=ur.role_id
                    WHERE u.user_id IN :ids ORDER BY u.username""").bindparams(
                bindparam("ids", expanding=True)
            ),
            {"ids": USER_IDS},
        )).mappings().all()
    print(*(dict(row) for row in rows), sep="\n")
    print(*(dict(user) for user in users), sep="\n")
    if not rows and not users:
        print("History demo data is not seeded.")


async def main(command: str) -> None:
    try:
        await {"seed": seed, "status": status, "reset": reset}[command]()
    finally:
        if engine is not None:
            await engine.dispose()


def run_cli(command: str) -> None:
    # psycopg's async connection is incompatible with Windows' default
    # ProactorEventLoop.  Keep this CLI aligned with scripts/run_local_api.py
    # so ``python -m app.history.demo ...`` works from PowerShell as documented.
    if sys.platform == "win32":
        asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())
    asyncio.run(main(command))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Manage Case-history demo data")
    parser.add_argument("command", choices=("seed", "status", "reset"))
    run_cli(parser.parse_args().command)
