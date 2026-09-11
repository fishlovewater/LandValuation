from io import BytesIO

import pytest
from docx import Document

from app.core.document_text_preview import (
    DocumentTextPreviewTooLargeError,
    build_docx_text_preview,
    preview_storage_docx,
)


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("土地估價報告", level=1)
    document.add_paragraph("案件編號 NB-2026-0001")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "欄位"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "估價基準日"
    table.cell(1, 1).text = "2026-09-11"

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_build_docx_text_preview_includes_paragraphs_and_tables() -> None:
    preview = build_docx_text_preview(_docx_bytes())

    assert preview.kind == "text"
    assert "土地估價報告" in preview.text
    assert "案件編號 NB-2026-0001" in preview.text
    assert "[表格]" in preview.text
    assert "欄位 | 值" in preview.text
    assert "估價基準日 | 2026-09-11" in preview.text
    assert preview.truncated is False


def test_build_docx_text_preview_reports_character_truncation() -> None:
    preview = build_docx_text_preview(_docx_bytes(), max_characters=12)

    assert len(preview.text) <= 12
    assert preview.truncated is True


class _Response:
    def __init__(self, content: bytes):
        self.content = content
        self.closed = False
        self.released = False

    def read(self, amount: int) -> bytes:
        return self.content[:amount]

    def close(self) -> None:
        self.closed = True

    def release_conn(self) -> None:
        self.released = True


class _Storage:
    def __init__(self, response: _Response):
        self.response = response

    async def download(self, object_key: str) -> _Response:
        assert object_key == "cases/demo/source.docx"
        return self.response


@pytest.mark.asyncio
async def test_preview_storage_docx_stops_reading_after_configured_limit() -> None:
    response = _Response(b"x" * 17)

    with pytest.raises(DocumentTextPreviewTooLargeError):
        await preview_storage_docx(
            _Storage(response),
            "cases/demo/source.docx",
            max_bytes=16,
        )

    assert response.closed is True
    assert response.released is True
