from io import BytesIO

from openpyxl import load_workbook

from app.valuation.report_packages.formal_xlsx_builder import (
    SHEET_F02,
    SHEET_F02_RF,
    SHEET_S01,
    SHEET_SUMMARY,
    build_formal_report_xlsx,
)


def _sample_data():
    return {
        "version_no": 4,
        "context": {
            "case_id": "11111111-1111-4111-8111-111111111111",
            "case_no": "NB-2026-0001",
            "case_title": "板橋區示範查估案件",
            "valuation_base_date": "2026-09-12",
            "city_code": "65000000",
            "district_code": "65000010",
            "land_use_type": "COMMERCIAL",
            "parcels": [
                {
                    "parcel_id": "22222222-2222-4222-8222-222222222222",
                    "district_code": "65000010",
                    "section_name": "文化段",
                    "subsection_name": "一小段",
                    "land_no": "100-1",
                    "area_sqm": "325.50",
                    "land_use_zone": "商業區",
                    "designated_use": None,
                }
            ],
            "benchmark_lands": [
                {
                    "benchmark_land_id": "33333333-3333-4333-8333-333333333333",
                    "parcel_id": "22222222-2222-4222-8222-222222222222",
                    "benchmark_land_no": "B-001",
                    "price_zone_no": "Z-10",
                }
            ],
        },
        "s01": {
            "district_name": "板橋區文化段",
            "district_boundary": "文化路周邊",
            "survey_date": "2026-09-10",
            "urban_plan_status": "都市計畫內",
            "land_use_zone": "商業區",
            "building_coverage_rate": "0.80",
            "floor_area_ratio": "4.60",
            "prohibited_building": "無",
            "restricted_building": "無",
            "main_road_name": "文化路",
            "main_road_width_m": "20.00",
            "average_road_width_m": "12.00",
            "observations": [
                {
                    "item_code": "mass_transit_proximity",
                    "raw_value": "近捷運站",
                    "facility_name": "板橋站",
                    "walking_distance_m": "450.00",
                    "source_notes": "現勘及路線資料確認",
                    "confirmed_by_user": True,
                }
            ],
            "notes": "",
            "site_opinion": "商業機能完整",
            "handler_name": "王承辦",
            "section_head_name": "李課長",
            "director_name": "陳主管",
            "appraiser_name": "林估價師",
        },
        "f02_rf": {
            "factor_rows": [
                {
                    "factor_code": "mass_transit_proximity",
                    "benchmark_reported_level": "佳",
                    "benchmark_confirmed_level": "佳",
                    "source_notes": "正式規則資料",
                    "confirmed_by_user": True,
                    "targets": [
                        {
                            "display_order": 1,
                            "reported_level": "普通",
                            "confirmed_level": "普通",
                            "calculated_adjustment_rate": "0.030000",
                            "source_notes": "人工確認",
                            "confirmed_by_user": True,
                        }
                    ],
                }
            ],
            "other_influences": "無",
            "notes": "",
            "appraiser_name": "林估價師",
            "calculation_status": "CALCULATED",
            "calculated_at": "2026-09-12T10:00:00+08:00",
        },
        "f02": {
            "comparison_workflow_enabled": True,
            "comparison_targets": [
                {
                    "display_order": 1,
                    "normal_unit_price_snapshot": "120000.00",
                    "transaction_date_snapshot": "2026-08-01",
                    "time_adjustment_rate": "0.010000",
                    "date_adjusted_price": "121200.00",
                    "regional_adjustment_rate": "0.030000",
                    "regional_adjusted_price": "124836.00",
                    "individual_adjustment_rate": "0.020000",
                    "trial_price": "127332.72",
                    "weight": "1.000000",
                    "weight_reason": "代表性最佳",
                    "individual_condition_notes": "臨路條件相近",
                    "individual_factors": [
                        {
                            "factor_code": "individual_area",
                            "benchmark_reported_level": "相近",
                            "benchmark_confirmed_level": "相近",
                            "comparable_reported_level": "略小",
                            "comparable_confirmed_level": "略小",
                            "calculated_adjustment_rate": "0.020000",
                            "source_notes": "面積比較",
                            "confirmed_by_user": True,
                        }
                    ],
                }
            ],
            "benchmark_comparison_price": "127333.00",
            "benchmark_notes": "採比較標的一",
            "notes": "",
            "handler_name": "王承辦",
            "section_head_name": "李課長",
            "director_name": "陳主管",
            "appraiser_name": "林估價師",
            "calculation_status": "CALCULATED",
            "calculated_at": "2026-09-12T10:00:00+08:00",
        },
    }


def test_formal_xlsx_uses_readable_labels_and_final_data_only():
    payload = build_formal_report_xlsx(_sample_data())
    workbook = load_workbook(BytesIO(payload), data_only=False)

    assert workbook.sheetnames == [SHEET_SUMMARY, SHEET_S01, SHEET_F02_RF, SHEET_F02]
    summary = workbook[SHEET_SUMMARY]
    summary_values = [cell.value for row in summary.iter_rows() for cell in row if cell.value]
    assert "板橋區" in summary_values
    assert "商業用地" in summary_values
    assert any("正式送審主要文件仍為完整送審 PDF" in str(value) for value in summary_values)

    all_text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "接近大型車站之程度" in all_text
    assert "面積" in all_text
    assert "mass_transit_proximity" not in all_text
    assert "individual_area" not in all_text
    assert "11111111-1111-4111-8111-111111111111" not in all_text
    assert "33333333-3333-4333-8333-333333333333" not in all_text
