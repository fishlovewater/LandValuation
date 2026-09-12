from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.valuation.report_packages.template_exports import (
    F04_FACTOR_ROWS,
    TEMPLATE_DIR,
    build_template_export_xlsx,
)


def _workbook(content: bytes):
    return load_workbook(BytesIO(content), data_only=False)


def test_official_workbook_already_contains_table_6_and_table_14() -> None:
    wb = load_workbook(
        TEMPLATE_DIR / "table5_residential_regional_factors.xlsx",
        read_only=True,
        data_only=False,
    )
    assert "表5-1區域因素明細表(住)" in wb.sheetnames
    assert "102表6宗地市價估計表格式" in wb.sheetnames
    assert "114表14比準地地價估計表" in wb.sheetnames
    assert "表5-2區域因素明細表(商)" not in wb.sheetnames


def test_commercial_f02_rf_export_fails_closed_without_official_template() -> None:
    with pytest.raises(FileNotFoundError, match="不會再由住宅表5-1自行改造"):
        build_template_export_xlsx(
            code="F02-RF",
            case={"case_no": "CASE-001", "land_use_type": "COMMERCIAL"},
            pages={"F02-RF": {}, "F02": {}},
        )


def test_f03_export_fills_table_14_and_clears_unused_sample_rows() -> None:
    content = build_template_export_xlsx(
        code="F03",
        case={"case_no": "CASE-014", "valuation_base_date": "2026-09-13"},
        pages={
            "F03": {
                "rows": [
                    {
                        "price_zone_no": "P001-00",
                        "benchmark_land_no": "1",
                        "district_name": "樹林區",
                        "section_subsection_name": "樹德段",
                        "land_no": "284",
                        "comparison_price": "137925",
                        "comparison_weight": "1",
                        "income_price": None,
                        "income_weight": "0",
                        "benchmark_land_price": "138000",
                        "decision_reason": "依比較法決定",
                    }
                ],
                "footer": {
                    "filled_date": "2026-09-13",
                    "handler_name": "承辦甲",
                    "section_head_name": "課長乙",
                    "director_name": "主任丙",
                    "appraiser_name": "估價師丁",
                },
            }
        },
    )
    wb = _workbook(content)
    ws = wb["114表14比準地地價估計表"]
    assert ws["B2"].value == "CASE-014"
    assert ws["I2"].value == "2026-09-13"
    assert ws["A5"].value == "P001-00"
    assert ws["B5"].value == "1"
    assert ws["C5"].value == "樹林區"
    assert ws["D5"].value == "樹德段"
    assert ws["E5"].value == "284"
    assert ws["F5"].value == 137925
    assert ws["J5"].value == 138000
    assert ws["K5"].value == "依比較法決定"
    assert ws["A6"].value is None
    assert "承辦甲" in ws["A10"].value
    assert ws["J11"].value == "估價師丁"
    assert "AI欄位對應" not in wb.sheetnames


def test_f04_export_fills_table_6_factor_rows_and_prices() -> None:
    parcel_id = str(uuid4())
    factors = [
        {
            "factor_code": code,
            "condition_value": f"條件-{index}",
            "difference_rate": "0.01",
            "source_notes": "正式規則",
            "confirmed_by_user": True,
        }
        for index, code in enumerate(F04_FACTOR_ROWS, start=1)
    ]
    content = build_template_export_xlsx(
        code="F04",
        case={"case_no": "CASE-006"},
        pages={
            "F04": {
                "benchmark_land_price": "100000",
                "benchmark": {
                    "benchmark_land_no": "B-01",
                    "parcel_display": "樹德段 284地號",
                },
                "benchmark_factor_rows": [
                    {
                        "factor_code": code,
                        "condition_value": f"比準-{index}",
                    }
                    for index, code in enumerate(F04_FACTOR_ROWS, start=1)
                ],
                "parcel_rows": [
                    {
                        "parcel_id": parcel_id,
                        "parcel_serial": 1,
                        "parcel_display": "樹德段 285地號",
                        "parcel_adjustment_rate": "0.20",
                        "parcel_unit_price": "120000",
                        "parcel_total_value": "12000000",
                        "factor_rows": factors,
                        "notes": "宗地備註",
                    }
                ],
                "calculation_snapshot": {
                    "rows": [
                        {
                            "parcel_id": parcel_id,
                            "trial_unit_price": "119999",
                        }
                    ]
                },
                "notes": "案件備註",
                "footer": {
                    "filled_date": "2026-09-13",
                    "handler_name": "承辦甲",
                    "section_head_name": "課長乙",
                    "director_name": "主任丙",
                    "appraiser_name": "估價師丁",
                },
            }
        },
    )
    wb = _workbook(content)
    ws = wb["102表6宗地市價估計表格式"]
    assert ws["O1"].value == "CASE-006"
    assert ws["E2"].value == "B-01"
    assert ws["I2"].value == 1
    assert ws["C4"].value == "樹德段 284地號"
    assert ws["F4"].value == "樹德段 285地號"
    assert ws["C6"].value == "比準-1"
    assert ws["F6"].value == "條件-1"
    assert ws["I6"].value == 0.01
    assert ws["F26"].value == 0.2
    assert ws["F27"].value == 119999
    assert ws["F28"].value == 120000
    assert ws["F29"].value == "宗地備註"
    assert ws["C30"].value == "案件備註"
    assert "承辦甲" in ws["A31"].value
    assert ws["M32"].value == "估價師丁"
    assert "AI欄位對應" not in wb.sheetnames
