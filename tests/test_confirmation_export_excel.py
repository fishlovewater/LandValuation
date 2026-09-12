from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace

import openpyxl

from app.valuation.automation.confirmation_export_excel import (
    build_confirmation_export_xlsx,
    get_field_label_zh,
)


def test_get_field_label_zh_returns_chinese_names() -> None:
    assert get_field_label_zh("S01", "administrative_area") == "行政區"
    assert get_field_label_zh("F01", "transaction_date") == "交易日期"
    assert get_field_label_zh("F03", "benchmark_land_no") == "比準地地號"
    assert get_field_label_zh("F02-RF", "building_coverage_rate") == "建蔽率"
    # Verify Item 8 fix for drainage_level
    assert get_field_label_zh("S01", "drainage_level") == "保（排）水之良否"
    assert get_field_label_zh("F02-RF", "drainage") == "排水之良否"
    assert get_field_label_zh("F02-RF", "vacancy_rate") == "店舖之毗連狀態"
    assert get_field_label_zh("F02-RF", "funeral_facility") == "殯葬設施之有無及接近程度"
    # Verify Item 1 physical factors
    assert get_field_label_zh("F02", "individual_area") == "面積 (M²)"
    assert get_field_label_zh("F02", "individual_width") == "寬度 (M)"
    assert get_field_label_zh("F02", "individual_depth") == "深度 (M)"
    assert get_field_label_zh("F02", "individual_shape") == "形狀"
    assert get_field_label_zh("F02", "individual_street_frontage") == "臨街情形"


def test_confirmation_export_xlsx_lists_all_fields_with_chinese_labels() -> None:
    xlsx_bytes = build_confirmation_export_xlsx(
        case_no="QA-CASE-001",
        case_title="確認結果測試",
        generated_at=datetime(2026, 8, 30, 9, 0, tzinfo=UTC),
        candidates=[
            SimpleNamespace(
                form_code="F02",
                field_name="valuation_base_date",
                field_status="CONFIRMED",
                confirmed_value="2025-05-28",
                extracted_value="114年5月28日",
                source_text="買賣日期：114年5月28日",
                page_number=1,
                analysis_provider="CODEX",
                confidence_score=0.95,
                created_at=datetime(2026, 8, 30, 9, 0, tzinfo=UTC),
            ),
            SimpleNamespace(
                form_code="F03",
                field_name="benchmark_land_no",
                field_status="REJECTED",
                confirmed_value=None,
                extracted_value="宗地流水號",
                source_text="此值不是比準地編號。",
                page_number=2,
                analysis_provider="OCR",
                confidence_score=0.80,
                created_at=datetime(2026, 8, 30, 9, 0, tzinfo=UTC),
            ),
        ],
    )

    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    assert "欄位確認結果" in wb.sheetnames
    ws = wb["欄位確認結果"]

    assert ws.cell(row=2, column=1).value == "案件：QA-CASE-001 確認結果測試"

    # Header is at row 5
    headers = [ws.cell(row=5, column=col).value for col in range(1, 15)]
    assert headers == [
        "表單",
        "欄位代碼",
        "候選群組",
        "不同候選答案數",
        "欄位中文名稱",
        "狀態",
        "確認/修正值",
        "擷取原始值",
        "來源證據原文",
        "頁碼",
        "分析來源",
        "信心度",
        "建立時間",
        "計算／缺值備註",
    ]

    row_count = ws.max_row
    assert row_count > 10

    found_f03_bm = False
    found_s01_drainage = False

    for row_idx in range(6, row_count + 1):
        form_code = ws.cell(row=row_idx, column=1).value
        field_name = ws.cell(row=row_idx, column=2).value
        label_zh = ws.cell(row=row_idx, column=5).value
        status = ws.cell(row=row_idx, column=6).value

        if form_code == "F03" and field_name == "benchmark_land_no":
            found_f03_bm = True
            assert label_zh == "比準地地號"
            assert status == "REJECTED"
        elif form_code == "S01" and field_name == "drainage_level":
            found_s01_drainage = True
            assert label_zh == "保（排）水之良否"

    assert found_f03_bm
    assert found_s01_drainage


def test_confirmation_export_calculates_only_when_no_direct_answer() -> None:
    xlsx_bytes = build_confirmation_export_xlsx(
        case_no="QA-CASE-002",
        case_title="formula fallback",
        generated_at=datetime(2026, 9, 8, 9, 0, tzinfo=UTC),
        candidates=[
            SimpleNamespace(
                form_code="F02",
                field_name="normal_land_unit_price",
                field_status="CONFIRMED",
                confirmed_value="100",
                extracted_value="100",
                source_text="normal price",
                page_number=1,
                analysis_provider="CODEX",
                confidence_score=0.95,
                created_at=datetime(2026, 9, 8, 9, 0, tzinfo=UTC),
            ),
            SimpleNamespace(
                form_code="F02",
                field_name="date_adjustment_rate",
                field_status="CONFIRMED",
                confirmed_value="2%",
                extracted_value="2%",
                source_text="date rate",
                page_number=1,
                analysis_provider="CODEX",
                confidence_score=0.95,
                created_at=datetime(2026, 9, 8, 9, 0, tzinfo=UTC),
            ),
        ],
    )

    ws = openpyxl.load_workbook(BytesIO(xlsx_bytes)).active
    rows = [
        [ws.cell(row=row_idx, column=column).value for column in range(1, 15)]
        for row_idx in range(6, ws.max_row + 1)
    ]
    adjusted = next(
        row for row in rows
        if row[0] == "F02" and row[1] == "adjusted_unit_price_display"
    )
    assert adjusted[5] == "FORMULA_CALCULATED"
    assert adjusted[6] == "102"
    assert adjusted[10] == "FORMULA"

    trial = next(
        row for row in rows
        if row[0] == "F02" and row[1] == "trial_price"
    )
    assert trial[5] == "FORMULA_INPUT_MISSING"


def test_confirmation_export_marks_multiple_answers_for_one_field() -> None:
    xlsx_bytes = build_confirmation_export_xlsx(
        case_no="QA-CASE-003",
        case_title="multiple answers",
        generated_at=datetime(2026, 9, 8, 9, 0, tzinfo=UTC),
        candidates=[
            SimpleNamespace(
                form_code="F03",
                field_name="benchmark_land_no",
                field_status="CONFIRMED",
                confirmed_value="P001-00",
                extracted_value="P001-00",
                source_text="document A",
                page_number=1,
                analysis_provider="OCR",
                confidence_score=0.9,
                created_at=datetime(2026, 9, 8, 9, 0, tzinfo=UTC),
            ),
            SimpleNamespace(
                form_code="F03",
                field_name="benchmark_land_no",
                field_status="REJECTED",
                confirmed_value=None,
                extracted_value="P002-00",
                source_text="document B",
                page_number=1,
                analysis_provider="OCR",
                confidence_score=0.8,
                created_at=datetime(2026, 9, 8, 9, 0, tzinfo=UTC),
            ),
        ],
    )

    ws = openpyxl.load_workbook(BytesIO(xlsx_bytes)).active
    rows = [
        [ws.cell(row=row_idx, column=column).value for column in range(1, 15)]
        for row_idx in range(6, ws.max_row + 1)
        if ws.cell(row=row_idx, column=1).value == "F03"
        and ws.cell(row=row_idx, column=2).value == "benchmark_land_no"
    ]

    assert len(rows) == 2
    assert {row[2] for row in rows} == {"F03.benchmark_land_no"}
    assert {row[3] for row in rows} == {2}
    assert all("請只確認一筆正確答案" in row[13] for row in rows)


def test_confirmation_export_includes_loaded_manual_values() -> None:
    xlsx_bytes = build_confirmation_export_xlsx(
        case_no="QA-CASE-MANUAL",
        case_title="manual values",
        generated_at=datetime(2026, 9, 10, 9, 0, tzinfo=UTC),
        candidates=[],
        manual_values={"S01": {"administrative_area": "65000010"}},
    )

    ws = openpyxl.load_workbook(BytesIO(xlsx_bytes)).active
    rows = [
        [ws.cell(row=row_idx, column=column).value for column in range(1, 15)]
        for row_idx in range(6, ws.max_row + 1)
    ]
    manual = next(
        row for row in rows
        if row[0] == "S01" and row[1] == "administrative_area"
    )
    assert manual[4] == "行政區"
    assert manual[5] == "MANUAL_CONFIRMED"
    assert manual[6] == "65000010"
    assert manual[10] == "MANUAL"
